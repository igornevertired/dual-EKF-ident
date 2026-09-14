#!/usr/bin/env python3
"""DOCX: цикл моделирования, модели и формулы ОФК-1 / ОФК-2 (OMML)."""

from __future__ import annotations

import sys
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml.ns import qn

_HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(_HERE))

from build_chapter1_docx import (  # noqa: E402
    D,
    E,
    add_body,
    add_heading,
    add_table,
    set_paragraph_format,
    set_run_font,
)

OUT = _HERE / "Архитектура_моделирования_ОФК1_ОФК2.docx"
PDF = _HERE / "architecture_ofk1_ofk2.pdf"
FIG = _HERE / "_arch_cycle.png"


def _render_arch_png() -> Path:
    if not PDF.exists():
        raise SystemExit(f"Нет {PDF}. Сначала: python docs/build_architecture_pdf.py")
    doc = fitz.open(PDF)
    pix = doc[1].get_pixmap(matrix=fitz.Matrix(2.0, 2.0))
    pix.save(FIG)
    return FIG


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, align="center", space_after=12)
    run = p.add_run(text)
    set_run_font(run, size=12, bold=True)


def add_figure(doc: Document, path: Path, width_cm: float = 16.0) -> None:
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, align="center", space_after=4)
    p.add_run().add_picture(str(path), width=Cm(width_cm))


def build() -> Path:
    _render_arch_png()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, align="center", space_after=6)
    run = p.add_run(
        "Цикл моделирования бортового навигационного комплекса и идентификации "
        "короткопериодических коэффициентов"
    )
    set_run_font(run, size=16, bold=True)

    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, align="center", space_after=18)
    run = p.add_run(
        "Описание по факту реализации: run_full_sim.py → bins_gnss_simulation.run_simulation"
    )
    set_run_font(run, size=12)

    add_heading(doc, "1. Назначение и структура контура", 1)
    add_body(
        doc,
        "В одном численном прогоне одновременно воспроизводятся: нелинейная динамика летательного "
        "аппарата (модель FX1), ошибки инерциальных датчиков, механизация бесплатформенной инерциальной "
        "навигационной системы (БИНС), упрощённые измерения ГНСС, оптимальный фильтр Калмана ошибок "
        "навигации (ОФК-1) и идентификатор короткопериодических продольных коэффициентов (ОФК-2). "
        "ОФК-1 не оценивает аэродинамику. ОФК-2 не оценивает координаты и скорость. Связь односторонняя: "
        "ОФК-1 корректирует навигационный вектор БИНС, после чего из скорректированного решения собираются "
        "угол атаки, скорость и тангаж для ОФК-2.",
    )
    add_body(
        doc,
        "Сценарий по умолчанию (" + E(r"t_{\mathrm{model}}=60") + " с): шаг интегрирования "
        + E(r"\Delta t=10^{-3}")
        + " с; такт ГНСС и ОФК-1 "
        + E(r"\Delta t_{\mathrm{GNSS}}=0{,}1")
        + " с (10 Гц); такт ОФК-2 "
        + E(r"\Delta t_{2}=0{,}02")
        + " с (50 Гц); балансировка горизонтального полёта "
        + E(r"H=500")
        + " м, "
        + E(r"V=80")
        + " м/с; манёвр руля высоты — doublet амплитуды "
        + E(r"\pm 5^{\circ}")
        + " (канал "
        + E(r"\delta V")
        + " автопилота на интервале манёвра заморожен). Автопилот вызывается каждые 0,01 с.",
    )

    add_figure(doc, FIG, 16.5)
    add_caption(doc, "Рисунок 1 — Полный цикл моделирования с двумя фильтрами")

    add_heading(doc, "2. Порядок вычислений на шаге интегрирования", 1)
    add_body(
        doc,
        "На шаге с индексом "
        + E("i")
        + " при "
        + E(r"t=(i+1)\Delta t")
        + " выполняется фиксированная последовательность.",
    )
    add_body(
        doc,
        "1. Интегрирование FX1 методом Рунге–Кутты 4-го порядка: состояние "
        + E(r"x\in\mathbb{R}^{25}")
        + ", идеальные показания "
        + E(r"\omega_{bi}^{b}")
        + " и кажущееся ускорение "
        + E(r"a_{f}^{b}")
        + " в связанной системе координат.",
    )
    add_body(
        doc,
        "2. Автопилот (каждые 0,01 с) формирует управление "
        + E("u")
        + ", включая doublet руля высоты.",
    )
    add_body(
        doc,
        "3. Генератор ошибок приборов (InsErrorGen) искажает идеальные выходы датчиков.",
    )
    add_body(
        doc,
        D(
            r"\omega_{m}=\omega_{bi}^{b}+b_{\omega}+n_{\omega},\qquad "
            r"a_{m}=a_{f}^{b}+b_{a}+n_{a}"
        ),
    )
    add_body(
        doc,
        "Смещения задаются один раз при создании генератора (seed 42): "
        + E(r"b_{\omega}\sim\mathcal{N}(0,\,0{,}5^{\circ}/\mathrm{ч})")
        + " по трём осям, "
        + E(r"b_{a}\sim\mathcal{N}(0,\,5\cdot 10^{-4}\ \mathrm{м/с}^{2})")
        + ". Белый шум на каждом шаге: "
        + E(r"\sigma_{\omega}=10^{-5}")
        + " (вход модели угловой скорости), "
        + E(r"\sigma_{a}=5\cdot 10^{-5}")
        + " м/с². Генератор ГНСС — отдельный RNG (seed 123).",
    )
    add_body(
        doc,
        "4. Механизация БИНС (bins_step): сначала навигационный канал по "
        + E(r"a_{m}")
        + ", затем ориентация по "
        + E(r"\omega_{m}")
        + ". Навигационный вектор "
        + E(r"n_{p}=[V_{N},\,V_{h},\,V_{E},\,h,\,\varphi,\,\lambda]^{\mathsf{T}}")
        + ".",
    )
    add_body(
        doc,
        "Ускорение в географической СК и производные навигационного вектора:",
    )
    add_body(
        doc,
        D(
            r"a_{\mathrm{nav}}=C_{bn}a_{m},\qquad "
            r"\dot{v}^{n}=a_{\mathrm{nav}}+a_{p}+a_{k}+a_{t}+g"
        ),
    )
    add_body(
        doc,
        "где "
        + E(r"a_{p}")
        + " — переносное (центростремительное) ускорение от вращения Земли, "
        + E(r"a_{k}")
        + " — кориолисово, "
        + E(r"a_{t}")
        + " — от вращения географической СК, "
        + E("g")
        + " — гравитация модели Земли. Кинематика координат:",
    )
    add_body(
        doc,
        D(
            r"\dot{h}=V_{h},\qquad "
            r"\dot{\varphi}=\frac{V_{N}}{R_{N}},\qquad "
            r"\dot{\lambda}=\frac{V_{E}}{R_{E}\cos\varphi}"
        ),
    )
    add_body(
        doc,
        "Ориентация: "
        + E(r"\omega_{bnb}^{b}=\omega_{m}-\omega_{ni}^{b}")
        + ", интегрирование кватерниона "
        + E(r"q")
        + " с нормировочной поправкой "
        + E(r"(1-\|q\|^{2})")
        + ", "
        + E(r"C_{bn}")
        + " восстанавливается из кватерниона. ОФК-1 не корректирует "
        + E("q")
        + " и "
        + E(r"C_{bn}")
        + ".",
    )
    add_body(
        doc,
        "5. Если "
        + E(r"(i+1)")
        + " кратно 100: измерение ГНСС, ОФК-1, обратная связь в "
        + E(r"n_{p}")
        + ". На первом такте скорости и высота БИНС выравниваются по ГНСС.",
    )
    add_body(
        doc,
        "6. Если уже был хотя бы один такт ОФК-1 и "
        + E(r"(i+1)")
        + " кратно 20: шаг ОФК-2 и параллельный шаг legacy Hoff. Истина FX1 в оба фильтра не подаётся; "
        "она используется для логов и для численного якобиана эталона ОФК-2.",
    )
    add_body(
        doc,
        "Следствие тактовой сетки. ОФК-2 работает на 50 Гц, ГНСС — на 10 Гц. Между тактами ОФК-1 величины "
        + E(r"\alpha,\,V,\,\theta")
        + " меняются только механизацией БИНС (ошибки ДУС/ДЛУ). Угловая скорость тангажа и нормальная "
        "перегрузка для ОФК-2 снимаются с IMU каждый такт 50 Гц.",
    )

    add_heading(doc, "3. Модель ГНСС", 1)
    add_body(
        doc,
        "Псевдодальности не моделируются. К истинным "
        + E(r"\varphi,\,\lambda,\,h,\,V_{N},\,V_{E},\,V_{h}")
        + " из FX1 добавляется независимый гауссов шум:",
    )
    add_body(
        doc,
        D(r"y_{\mathrm{GNSS}}=y_{\mathrm{true}}+\sigma\odot\xi,\qquad \xi\sim\mathcal{N}(0,I_{6})"),
    )
    add_table(
        doc,
        ["Канал", "СКО по умолчанию"],
        [
            ["широта φ", "5·10⁻⁶ рад"],
            ["долгота λ", "5·10⁻⁶ рад"],
            ["высота h", "1 м"],
            ["Vn, Ve, Vh", "0,2 м/с"],
        ],
    )

    add_heading(doc, "4. ОФК-1: оценка ошибок БИНС", 1)
    add_body(
        doc,
        "Тип: error-state фильтр Калмана, loosely coupled. Состояние "
        + E(r"x_{\mathrm{I}}\in\mathbb{R}^{15}")
        + ": первые 13 компонент — порт модели ошибок БИНС (BINS_OFK_2ch.m), индексы 13 и 14 — "
        + E(r"\delta V_{h}")
        + " и "
        + E(r"\delta h")
        + ". Измерение — шесть каналов ГНСС.",
    )
    add_body(
        doc,
        "Инновация (порядок строк матрицы "
        + E("H")
        + "):",
    )
    add_body(
        doc,
        D(
            r"z=\begin{bmatrix}"
            r"\varphi_{\mathrm{BINS}}-\varphi_{\mathrm{GNSS}}\\"
            r"\lambda_{\mathrm{BINS}}-\lambda_{\mathrm{GNSS}}\\"
            r"V_{N,\mathrm{BINS}}-V_{N,\mathrm{GNSS}}\\"
            r"V_{E,\mathrm{BINS}}-V_{E,\mathrm{GNSS}}\\"
            r"V_{h,\mathrm{BINS}}-V_{h,\mathrm{GNSS}}\\"
            r"h_{\mathrm{BINS}}-h_{\mathrm{GNSS}}"
            r"\end{bmatrix}"
        ),
    )
    add_body(
        doc,
        "Линеаризованная модель ошибок на такте ГНСС:",
    )
    add_body(
        doc,
        D(r"\dot{x}=F x+G w,\qquad z=H x+\nu"),
    )
    add_body(
        doc,
        "Матрицы "
        + E("F,G,H,Q")
        + " строятся из "
        + E(r"a_{m}")
        + " (тело), "
        + E(r"C_{bn}")
        + " и "
        + E(r"n_{p}")
        + ". Для 13-мерного ядра "
        + E("H")
        + " выделяет "
        + E(r"\delta\varphi,\,\delta\lambda,\,\delta V_{N},\,\delta V_{E}")
        + "; расширение 15-го порядка добавляет "
        + E(r"H_{5,14}=1")
        + ", "
        + E(r"H_{6,15}=1")
        + " (индексация с единицы). Вертикальный канал:",
    )
    add_body(
        doc,
        D(r"\delta\dot{V}_{h}=\ldots,\qquad \delta\dot{h}=\delta V_{h}"),
    )
    add_body(
        doc,
        "без кросс-связей с горизонтальными скоростями (явное ограничение в коде, чтобы не портить "
        "горизонтальные каналы через ковариацию). Процессный шум "
        + E("Q")
        + " как в MATLAB:",
    )
    add_body(
        doc,
        D(r"Q=\mathrm{diag}\bigl((w_{q}^{\circ 2})/\Delta t_{\mathrm{GNSS}}\bigr)"),
    )
    add_body(
        doc,
        "с "
        + E(r"w_{q}")
        + " = (20·10⁻¹⁰, 20·10⁻¹⁰, 20·10⁻¹⁰, 25·10⁻⁷, 25·10⁻⁷, 25·10⁻⁷). "
        "Ковариация измерения "
        + E("R")
        + " — диагональ из "
        + E(r"\sigma_{\mathrm{GNSS}}^{2}")
        + " с нижним полом и перестановкой под порядок "
        + E("H")
        + ".",
    )
    add_body(doc, "Дискретизация предсказания (как OFK.m):")
    add_body(
        doc,
        D(
            r"\Phi=I+F\Delta t+\frac{(F\Delta t)^{2}}{2},\qquad "
            r"G_{\mathrm{rus}}=G\Delta t"
        ),
    )
    add_body(
        doc,
        D(
            r"\hat{x}^{-}=\Phi\hat{x}^{+},\qquad "
            r"P^{-}=\Phi P^{+}\Phi^{\mathsf{T}}+G_{\mathrm{rus}}Q G_{\mathrm{rus}}^{\mathsf{T}}"
        ),
    )
    add_body(doc, "Обновление:")
    add_body(
        doc,
        D(
            r"S=HP^{-}H^{\mathsf{T}}+R,\qquad "
            r"K=P^{-}H^{\mathsf{T}}S^{-1},\qquad "
            r"\hat{x}^{+}=\hat{x}^{-}+K(z-H\hat{x}^{-})"
        ),
    )
    add_body(
        doc,
        D(
            r"P^{+}=(I-KH)P^{-}(I-KH)^{\mathsf{T}}+KRK^{\mathsf{T}}"
        ),
    )
    add_body(
        doc,
        "Обратная связь только в "
        + E(r"n_{p}")
        + " (не в состояние FX1):",
    )
    add_body(
        doc,
        D(
            r"V_{N}\leftarrow V_{N}-\hat{x}_{4},\ "
            r"V_{E}\leftarrow V_{E}-\hat{x}_{5},\ "
            r"\varphi\leftarrow\varphi-\hat{x}_{6},\ "
            r"\lambda\leftarrow\lambda-\hat{x}_{7},\ "
            r"V_{h}\leftarrow V_{h}-\hat{x}_{14},\ "
            r"h\leftarrow h-\hat{x}_{15}"
        ),
    )
    add_body(
        doc,
        "Индексация в формуле — с единицы, как в тексте MATLAB-порядка; в коде Python это "
        "x[3], x[4], x[5], x[6], x[13], x[14]. После обратной связи "
        + E(r"\hat{x}_{\mathrm{I}}:=0")
        + " (error-state reset). Начальная "
        + E(r"P_{0}")
        + " диагональная: 10³ (первые 4), 10⁻² (следующие 6), 10⁻⁴ (остальные), "
        + E(r"P_{0}[13,13]=0{,}04")
        + ", "
        + E(r"P_{0}[14,14]=1")
        + ".",
    )

    add_heading(doc, "5. ОФК-2: идентификация L*, M*", 1)
    add_heading(doc, "5.1. Модель измерений", 2)
    add_body(
        doc,
        "Короткопериодическая продольная модель в приращениях от точки балансировки "
        + E(r"t=0")
        + " (trim):",
    )
    add_body(
        doc,
        D(
            r"\dot{\alpha}=L_{\alpha}\delta\alpha+L_{q}q+L_{\delta e}\delta\delta e"
            r"+L_{V}\delta V+L_{\theta}\delta\theta"
        ),
    )
    add_body(
        doc,
        D(
            r"\dot{q}=M_{\alpha}\delta\alpha+M_{q}q+M_{\delta e}\delta\delta e"
            r"+M_{V}\delta V+M_{\theta}\delta\theta"
        ),
    )
    add_body(
        doc,
        D(
            r"a_{z}\approx a_{z0}-\frac{V}{g}\Bigl("
            r"L_{\alpha}\delta\alpha+(L_{q}-1)q+L_{\delta e}\delta\delta e"
            r"+L_{V}\delta V+L_{\theta}\delta\theta\Bigr)"
        ),
    )
    add_body(
        doc,
        "Член "
        + E(r"(L_{q}-1)q")
        + " отделяет кинематический вклад "
        + E(r"\dot{\alpha}\approx q")
        + " от аэродинамической части в нормальной перегрузке. На балансировке численный якобиан FX1 даёт "
        + E(r"L_{q}\approx 1")
        + ".",
    )
    add_body(
        doc,
        "Состояние фильтра — вектор из десяти коэффициентов:",
    )
    add_body(
        doc,
        D(
            r"\theta=\bigl[L_{\alpha},\,L_{q},\,L_{\delta e},\,L_{V},\,L_{\theta},"
            r"\,M_{\alpha},\,M_{q},\,M_{\delta e},\,M_{V},\,M_{\theta}\bigr]^{\mathsf{T}}"
        ),
    )
    add_body(
        doc,
        "Регрессор "
        + E(r"\phi\in\mathbb{R}^{5}")
        + ":",
    )
    add_body(
        doc,
        D(
            r"\phi=\bigl[\delta\alpha,\,q,\,\delta\delta e,\,\delta V,\,\delta\theta\bigr]^{\mathsf{T}}"
        ),
    )
    add_body(
        doc,
        D(
            r"\delta\alpha=\alpha-\alpha_{0},\quad "
            r"\delta\delta e=\delta e-\delta e_{0},\quad "
            r"\delta V=V-V_{0},\quad "
            r"\delta\theta=\theta-\theta_{0}"
        ),
    )
    add_body(
        doc,
        "Источники сигналов. "
        + E(r"\alpha")
        + " и "
        + E("V")
        + " из навигации: скорость в связанной СК "
        + E(r"v^{b}=C_{bn}^{\mathsf{T}}[V_{N},V_{h},V_{E}]^{\mathsf{T}}")
        + ", "
        + E(r"\alpha=-\mathrm{arctan2}(v_{y}^{b},v_{x}^{b})")
        + ". Тангаж "
        + E(r"\theta")
        + " из "
        + E(r"C_{bn}")
        + " (c_ang). "
        + E(r"q=\omega_{m,z}")
        + ", "
        + E(r"a_{z}=a_{m,y}/g")
        + ". Руль "
        + E(r"\delta e")
        + " — состояние привода FX1 (градусы, перевод в радианы). "
        + E(r"a_{z0}")
        + " фиксируется на первом такте измерения "
        + E(r"a_{z}")
        + " и далее не обновляется.",
    )

    add_heading(doc, "5.2. Equation-error EKF", 2)
    add_body(
        doc,
        "Производные — конечные разности за шаг ОФК-2:",
    )
    add_body(
        doc,
        D(
            r"\dot{\alpha}_{\mathrm{fd}}=\frac{\delta\alpha_{k}-\delta\alpha_{k-1}}{\Delta t_{2}},\qquad "
            r"\dot{q}_{\mathrm{fd}}=\frac{q_{k}-q_{k-1}}{\Delta t_{2}}"
        ),
    )
    add_body(
        doc,
        "Вектор измерения "
        + E(r"Z=[\dot{\alpha}_{\mathrm{fd}},\,\dot{q}_{\mathrm{fd}},\,a_{z}]^{\mathsf{T}}")
        + ". Предсказание измерения линейно по "
        + E(r"\theta")
        + ". Матрица наблюдения "
        + E(r"H\in\mathbb{R}^{3\times 10}")
        + ":",
    )
    add_body(
        doc,
        D(
            r"H=\begin{bmatrix}"
            r"\phi^{\mathsf{T}} & 0_{1\times 5}\\"
            r"0_{1\times 5} & \phi^{\mathsf{T}}\\"
            r"k_{az}\phi^{\mathsf{T}} & 0_{1\times 5}"
            r"\end{bmatrix},\qquad "
            r"k_{az}=-V/g"
        ),
    )
    add_body(
        doc,
        "Динамика состояния: "
        + E(r"\dot{\theta}=0")
        + ", процессный шум "
        + E(r"Q=0")
        + ". Ковариация измерения:",
    )
    add_body(
        doc,
        D(
            r"R=\mathrm{diag}(0{,}025^{2},\,0{,}002^{2},\,0{,}023^{2})"
        ),
    )
    add_body(
        doc,
        "Шаг EKF выполняется только при наличии предыдущего такта для разностей и при возбуждении: "
        + E(r"|\delta\delta e|\ge 0{,}003")
        + " либо "
        + E(r"|q|\ge 0{,}005")
        + " либо "
        + E(r"|\delta\alpha|\ge 0{,}008")
        + " либо "
        + E(r"|\delta V|\ge 0{,}05")
        + " либо "
        + E(r"|\delta\theta|\ge 0{,}003")
        + " (радианы, м/с). Иначе обновления нет, "
        + E("P")
        + " не раздувается процессом (Q нулевой).",
    )
    add_body(
        doc,
        "Начальное условие: "
        + E(r"\hat{\theta}(0)=1{,}3\,\theta_{\mathrm{trim}}")
        + " (смещение +30 %). Диагональ "
        + E(r"P_{0}")
        + " задаётся так, чтобы 3σ покрывала это смещение, с абсолютным полом по каналам.",
    )

    add_heading(doc, "5.3. Эталон и метрики", 2)
    add_body(
        doc,
        "Эталон — численный якобиан FX1 (центральные разности по "
        + E(r"\alpha,\,q,\,\delta e,\,V,\,\theta")
        + "). "
        + E(r"\theta_{\mathrm{trim}}")
        + " считается один раз в "
        + E(r"(x_{0},u_{0})")
        + " и не меняется. "
        + E(r"\theta_{\mathrm{FX1}}(t)")
        + " пересчитывается каждый такт 50 Гц из текущего "
        + E(r"(x,u)")
        + " FX1. Фильтр эталон не использует (нет theory-anchor). На графиках ошибки:",
    )
    add_body(
        doc,
        D(
            r"\Delta_{\mathrm{trim}}(t)=\hat{\theta}(t)-\theta_{\mathrm{trim}},\qquad "
            r"\Delta_{\mathrm{inst}}(t)=\hat{\theta}(t)-\theta_{\mathrm{FX1}}(t)"
        ),
    )
    add_body(
        doc,
        "Чёрная пунктирная линия на графике старт/финиш — дрейф эталона "
        + E(r"\theta_{\mathrm{FX1}}(t)-\theta_{\mathrm{trim}}")
        + ". Оранжевая линия — нуль (совпадение с балансировкой).",
    )
    add_body(
        doc,
        "Коэффициенты "
        + E(r"L_{\theta},\,M_{\theta}")
        + " входят в состояние фильтра (структура уравнений), на балансировке их якобиан ≈ 0, наблюдаемость слабая. "
        "В отчёт выводятся восемь параметров: "
        + E(r"L_{\alpha},L_{q},L_{\delta e},L_{V},M_{\alpha},M_{q},M_{\delta e},M_{V}")
        + ".",
    )

    add_heading(doc, "5.4. Legacy Hoff (только сравнение)", 2)
    add_body(
        doc,
        "Параллельно на тех же "
        + E(r"\delta\alpha,\,q,\,\delta\delta e,\,a_{z}")
        + " работает 6-параметрическая модель:",
    )
    add_body(
        doc,
        D(
            r"\dot{\alpha}=L_{\alpha}\delta\alpha+L_{q}q+L_{\delta e}\delta\delta e"
        ),
    )
    add_body(
        doc,
        D(
            r"\dot{q}=M_{\alpha}\delta\alpha+M_{q}q+M_{\delta e}\delta\delta e"
        ),
    )
    add_body(
        doc,
        "На навигацию и на основную ОФК-2 не влияет. Нужна как база: без "
        + E(r"\delta V,\,\delta\theta")
        + " ошибки БНК по скорости и тангажу входят в оценки "
        + E(r"L_{\alpha},L_{q},M_{\alpha}")
        + ".",
    )

    add_heading(doc, "6. Файлы реализации", 1)
    add_table(
        doc,
        ["Блок", "Файл"],
        [
            ["Точка входа", "run_full_sim.py"],
            ["Цикл времени", "simulation/bins_gnss_simulation.py"],
            ["Балансировка, x0", "simulation/initsim.py"],
            ["Динамика ЛА", "dynamics/fx1.py"],
            ["Автопилот", "dynamics/autopilot_model.py"],
            ["Ошибки ДУС/ДЛУ", "sensors/imu_error_generator.py"],
            ["ГНСС", "sensors/core.py → gnss()"],
            ["Механизация БИНС", "navigation/bins_common.py"],
            ["ОФК-1", "filtering/loosely_coupled_ofk.py, bins_ofk_2ch.py"],
            ["ОФК-2", "filtering/ofk2_ekf.py"],
            ["Эталон якобиан", "filtering/ofk2_theory.py"],
            ["Hoff 6 коэфф.", "filtering/ofk2_ekf_legacy.py"],
            ["Графики", "output/full_sim_outputs.py"],
        ],
    )

    add_heading(doc, "7. Что реализовано и чего нет", 1)
    add_body(
        doc,
        "Реализовано: полный контур FX1 → датчики → БИНС → ГНСС → ОФК-1 → ОФК-2; расширенная "
        "10-параметрическая модель идентификации; параллельный Hoff; эталон якобианом FX1; графики "
        "навигации и идентификации. Есть прогон по амплитуде doublet (возбуждение), это не sweep ошибок БНС.",
    )
    add_body(
        doc,
        "Не реализовано: серия прогонов по уровню ошибок ГНСС/БИНС (множитель σ в run_simulation не задаётся). "
        "ОФК-1 не корректирует ориентацию. "
        + E(r"a_{z0}")
        + " заморожен. "
        + E("Q")
        + " ОФК-2 нулевой. "
        + E(r"L_{\theta},M_{\theta}")
        + " остаются в фильтре и исключены из отчётных графиков.",
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"DOCX: {path}")
