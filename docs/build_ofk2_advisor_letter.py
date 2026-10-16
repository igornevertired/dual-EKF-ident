"""Письмо научруку: ОФК-2 (вторая часть). Запуск: poetry run python docs/build_ofk2_advisor_letter.py"""

from __future__ import annotations

import re
from pathlib import Path

import latex2mathml.converter
import mathml2omml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
OUT = Path(__file__).resolve().parent / "ОФК-2_модели_для_научника.docx"
PLOTS = Path(__file__).resolve().parents[1] / "src" / "plots"


def latex_to_omath_xml(latex: str) -> str:
    mml = latex2mathml.converter.convert(latex)
    omml = mathml2omml.convert(mml)
    if "xmlns:m=" not in omml:
        omml = omml.replace("<m:oMath>", f'<m:oMath xmlns:m="{M_NS}">', 1)
    return omml


def add_omath(paragraph, latex: str) -> None:
    omml = latex_to_omath_xml(latex)
    omml = re.sub(r'\sxmlns:m="[^"]*"', "", omml)
    wrapped = (
        f'<m:oMath xmlns:m="{M_NS}">'
        + omml.replace("<m:oMath>", "").replace("</m:oMath>", "")
        + "</m:oMath>"
    )
    paragraph._p.append(parse_xml(wrapped))


def add_display(doc: Document, latex: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    add_omath(p, latex)


def add_runs_with_math(p, text: str) -> None:
    parts = re.split(r"(\$[^$]+\$)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("$") and part.endswith("$") and len(part) > 2:
            add_omath(p, part[1:-1])
        else:
            p.add_run(part)


def set_run_font(run, *, bold=False, size=12, name="Times New Roman"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    r = run._element
    rpr = r.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)


def heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    add_runs_with_math(p, text)
    for run in p.runs:
        set_run_font(run, bold=True, size=14 if level == 1 else 13)


def para(doc: Document, text: str, *, first_indent: bool = False) -> None:
    p = doc.add_paragraph()
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    add_runs_with_math(p, text)
    for run in p.runs:
        set_run_font(run)


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        add_runs_with_math(p, item)
        for run in p.runs:
            set_run_font(run)


def add_figure(doc: Document, path: Path, caption: str, *, width_in: float = 5.8) -> None:
    if not path.is_file():
        para(doc, f"[рисунок не найден: {path.name}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width_in))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    set_run_font(r, size=11)
    r.italic = True


def build() -> Path:
    doc = Document()
    style_doc(doc)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("ОФК-2: модель, алгоритм и текущие результаты")
    set_run_font(r, bold=True, size=16)

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run("после ОФК-1 → сигналы потребителя → три линейных ДУ → оценка L*, M*, X*")
    set_run_font(r, size=12)
    r.italic = True

    para(doc, "Добрый день.")
    para(
        doc,
        "Ниже — вторая часть контура: ОФК-2 так, как сейчас сделано в программе. "
        "Прошу посмотреть и подсказать, в какую сторону двигаться дальше "
        "(оставить три ДУ как есть / сузить набор оцениваемых коэффициентов / расширять модель).",
        first_indent=True,
    )

    heading(doc, "0. Место ОФК-2 в контуре")
    bullets(
        doc,
        [
            "ОФК-1 оценивает ошибки БИНС; в механизацию БИНС оценки не пишет.",
            "Потребителю отдаём: NP минус $\\delta$ ОФК-1, "
            "$\\boldsymbol{\\omega}_m-\\Delta\\hat{\\boldsymbol{\\omega}}$, "
            "$\\mathbf{a}_m-\\Delta\\hat{\\mathbf{a}}$, матрицу $C_b^n$ с БИНС (без поправки $\\hat{\\varphi}$).",
            "ОФК-2 — отдельный фильтр идентификации аэродинамических коэффициентов. "
            "На полёт и на БИНС не влияет. Частота та же, что у ОФК-1: 10 Гц.",
            "Манёвр: дублет руля высоты (возбуждение короткого периода).",
        ],
    )

    heading(doc, "1. Что приходит на вход ОФК-2")
    bullets(
        doc,
        [
            "$\\alpha$, $V$ — из NP потребителя и $C_b^n$ БИНС "
            "(проекция скорости в связанную СК, $\\alpha=-\\mathrm{atan2}(V_y,V_x)$);",
            "$q=\\omega_z$ — из ДУС потребителя;",
            "$\\delta_e$ — угол руля из модели самолёта (управляющий вход);",
            "$\\delta\\vartheta$ — отклонение тангажа от балансировки. "
            "Считаем согласованно с NP: $\\vartheta\\approx\\alpha+\\arcsin(V_h/V)$, "
            "а не сырой угол из $C_b^n$ (иначе невязка по тангажу ломает фильтр);",
            "точка балансировки $(\\alpha_0,q_0,V_0,\\vartheta_0,\\delta_{e0})$ — из trim FX1.",
        ],
    )
    para(doc, "Измерение фильтра:")
    add_display(doc, r"Z=[\alpha,\ q,\ V]^{\top}")

    heading(doc, "2. Модель ОФК-2 (три линейных ДУ)")
    para(
        doc,
        "Около балансировки отклонения "
        "$\\delta\\alpha=\\alpha-\\alpha_0$, $\\delta V=V-V_0$, "
        "$\\delta\\delta_e=\\delta_e-\\delta_{e0}$, $\\delta\\vartheta=\\vartheta-\\vartheta_0$. "
        "Регрессор:",
    )
    add_display(doc, r"\varphi=[\delta\alpha,\ q,\ \delta\delta_e,\ \delta V]^{\top}")
    para(doc, "Продольная линейная модель:")
    add_display(
        doc,
        r"\dot{\alpha}=L_\alpha\delta\alpha+L_q q+L_{\delta e}\delta\delta_e+L_V\delta V",
    )
    add_display(
        doc,
        r"\dot{q}=M_\alpha\delta\alpha+M_q q+M_{\delta e}\delta\delta_e+M_V\delta V",
    )
    add_display(
        doc,
        r"\dot{V}=X_\alpha\delta\alpha+X_q q+X_{\delta e}\delta\delta_e+X_V\delta V"
        r"+X_\vartheta\delta\vartheta",
    )
    bullets(
        doc,
        [
            "Коэффициенты $L_*$, $M_*$, $X_*$ — то, что оцениваем (12 чисел).",
            "$X_\\vartheta$ берём из якобиана нелинейной FX1 в trim ($\\approx -g$), "
            "в состояние фильтра не кладём: это известный вклад тяжести в $\\dot V$.",
            "Уравнений $\\dot\\vartheta$, $\\dot h$ в фильтре нет — только три ДУ выше.",
        ],
    )

    heading(doc, "3. Откуда эталон «истины» для сравнения")
    para(
        doc,
        "Эталон $L_*$, $M_*$, $X_*$ — не таблицы из паспорта, а численный якобиан полной "
        "нелинейной модели FX1 в точке балансировки: частные производные "
        "$(\\dot\\alpha,\\dot q,\\dot V)$ по $(\\alpha,q,\\delta_e,V,\\vartheta)$. "
        "На графиках ошибка = оценка фильтра минус этот trim-эталон.",
    )

    heading(doc, "4. Состояние и шаг EKF")
    para(doc, "Состояние (15 чисел):")
    add_display(
        doc,
        r"\mathbf{x}=["
        r"L_\alpha,L_q,L_{\delta e},L_V,"
        r"M_\alpha,M_q,M_{\delta e},M_V,"
        r"X_\alpha,X_q,X_{\delta e},X_V,"
        r"\delta\alpha,q,\delta V"
        r"]^{\top}",
    )
    bullets(
        doc,
        [
            "Прогноз: интегрируем три ДУ на шаге $dt=0.1$ с, коэффициенты считаем постоянными "
            "($\\dot L=\\dot M=\\dot X=0$).",
            "Коррекция по $Z=[\\alpha,q,V]$; $P$ — по формуле Джозефа.",
            "Старт коэффициентов: trim $\\times(1+\\varepsilon)$ (намеренный сдвиг), "
            "чтобы было видно сходимость.",
        ],
    )

    heading(doc, "5. Чем ОФК-2 отличается от FX1 (важное)")
    para(
        doc,
        "Данные полёта идут из полной нелинейной FX1: аэродинамика $C_x(C_y^2)$, "
        "плотность $\\rho(h)$, тяжесть через $C_{gb}(\\vartheta,\\gamma,\\psi)$, "
        "полный 6DOF, отдельные уравнения $\\vartheta$, $h$ и т.д.",
    )
    para(
        doc,
        "ОФК-2 оценивает коэффициенты упрощённой линейной модели около одной точки trim. "
        "На дублете самолёт уходит от балансировки: появляются эффекты, которых в регрессоре "
        "$\\varphi$ нет. Это не ошибка программирования фильтра, а ограничение модели.",
    )

    heading(doc, "6. Что получается сейчас")
    bullets(
        doc,
        [
            "Короткопериод по $\\alpha$, $q$ (в основном $L_\\alpha$, $L_q$, $M_\\alpha$, $M_q$, "
            "$M_{\\delta e}$) — фильтр ловит, ошибки ближе к полосе $\\pm 3\\sigma$.",
            "$L_V$, $M_V$ на дублете почти не возбуждаются (на $\\sim 2$ порядка меньше "
            "$L_\\alpha$, $M_\\alpha$) — оценить их по сути нельзя.",
            "$X_q\\approx 0$ в trim — оценивать нечего.",
            "$X_\\alpha$, $X_V$ стабильно выходят за $\\pm 3\\sigma$ при «уверенном» фильтре "
            "(типично $\\Delta X_\\alpha\\sim +1.4$ при MAE по всем 12 $\\sim 0.15$). "
            "Это похоже на смещение из-за несоответствия моделей, а не на шум измерений.",
        ],
    )

    add_figure(
        doc,
        PLOTS / "ofk2_error_vs_three_sigma_P.png",
        "Рис. 1. Ошибка параметров ОФК-2 (оценка − trim) и полосы ±3σ фильтра",
    )
    add_figure(
        doc,
        PLOTS / "ofk2_params_time.png",
        "Рис. 2. Оценки L*, M*, X* во времени против эталона trim",
        width_in=5.8,
    )

    heading(doc, "7. Вопрос: в какую сторону смотреть дальше?")
    para(doc, "Прошу указать предпочтительный вариант (или комбинацию):")
    bullets(
        doc,
        [
            "Оставить три ДУ и оценивать все 12 коэффициентов, в тексте ВКР явно "
            "зафиксировав ограничение линейности и слабую идентифицируемость $L_V$, $M_V$, $X_*$.",
            "Оценивать только короткопериод "
            "$(L_\\alpha,L_q,L_{\\delta e},M_\\alpha,M_q,M_{\\delta e})$, "
            "а $L_V$, $M_V$, $X_*$ держать константами с trim FX1 "
            "(те же три ДУ, но маска на оцениваемые).",
            "Расширять модель: уравнения $\\vartheta$, $h$ / нелинейные члены / "
            "меньший манёвр ближе к trim — чтобы лучше тянуть $X_*$.",
            "Другое — по Вашей пометке.",
        ],
    )

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(16)
    r = note.add_run(
        "Готов уточнить любой блок и пересобрать прогон под выбранный вариант."
    )
    set_run_font(r)
    r.italic = True

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"saved: {path}")
