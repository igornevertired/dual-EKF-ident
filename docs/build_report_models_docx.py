#!/usr/bin/env python3
"""«Модели и алгоритмы» — исходный текст автора, дополненный формулами и графиками.

Порядок сборки::

    python docs/build_report_figures.py    # окружение с matplotlib
    python docs/_report_stats.py           # то же окружение, числа прогона
    python docs/build_report_models_docx.py  # окружение с python-docx

Результат: ``docs/Отчет_модели_и_алгоритмы_дополненный.docx``
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Cm, Pt
from docx.oxml.ns import qn

_HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(_HERE))

from build_chapter1_docx import (  # noqa: E402
    D,
    E,
    add_body,
    add_heading,
    add_table,
    set_paragraph_format,
    set_run_font,
)

ROOT = _HERE.parent
PLOTS = ROOT / "src" / "plots"
ARCH = _HERE / "_report_arch.png"
TIMING = _HERE / "_report_timing.png"
STATS = _HERE / "_report_stats.json"
OUT = _HERE / "Отчет_модели_и_алгоритмы_дополненный.docx"


def _fmt(value: float, digits: int = 3) -> str:
    if value != 0 and (abs(value) < 1e-3 or abs(value) >= 1e5):
        return f"{value:.{digits}e}".replace(".", ",")
    return f"{value:.{digits}f}".replace(".", ",")


def _pct(value: float) -> str:
    return f"{value:.1f}".replace(".", ",")


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, align="center", space_after=14)
    run = p.add_run(text)
    set_run_font(run, size=12)


def add_table_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, align="left", space_after=4)
    run = p.add_run(text)
    set_run_font(run, size=12)


def add_figure(doc: Document, path: Path, width_cm: float = 16.0) -> None:
    if not path.exists():
        raise SystemExit(f"Нет рисунка {path}")
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, align="center", space_after=4)
    p.add_run().add_picture(str(path), width=Cm(width_cm))


def add_page_break(doc: Document) -> None:
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, space_after=0)
    p.add_run().add_break(WD_BREAK.PAGE)


# ======================================================================


def _s1(doc: Document) -> None:
    add_heading(doc, "1. Назначение контура", 1)
    add_body(
        doc,
        "В одном прогоне считаются: динамика самолёта, ошибки ДУС/ДЛУ, навигация БИНС, "
        "упрощённые измерения ГНСС, поправка навигации (ОФК-1) и идентификация коэффициентов "
        "короткопериодического движения (ОФК-2).",
    )
    add_body(
        doc,
        "ОФК-1 не оценивает аэродинамику. ОФК-2 не оценивает координаты и скорость. Связь "
        "односторонняя: ОФК-1 правит вектор БИНС, из этого решения для ОФК-2 собираются угол "
        "атаки, скорость и тангаж. Истина FX1 в фильтры не подаётся. Она нужна как эталон для "
        "графиков и для численного якобиана коэффициентов.",
    )
    add_body(
        doc,
        "Полная схема контура приведена на рисунке 1; такты подсистем и профиль манёвра — на "
        "рисунке 2, сводка тактов — в таблице 1.",
    )
    add_figure(doc, ARCH, 16.5)
    add_caption(doc, "Рисунок 1 — Схема контура моделирования")
    add_table_caption(doc, "Таблица 1 — Такты подсистем")
    add_table(
        doc,
        ["Подсистема", "Шаг", "Частота"],
        [
            ["FX1 + Рунге–Кутта, генератор ошибок ДУС/ДЛУ, механизация БИНС", "1·10⁻³ с", "1000 Гц"],
            ["Автопилот", "0,01 с", "100 Гц"],
            ["ОФК-2 (идентификация)", "0,02 с", "50 Гц"],
            ["ГНСС + ОФК-1 (навигация)", "0,1 с", "10 Гц"],
            ["Длительность прогона", "60 с", "—"],
        ],
    )
    add_figure(doc, TIMING, 16.0)
    add_caption(
        doc,
        "Рисунок 2 — Профиль отклонения руля высоты и такты подсистем внутри одного такта ГНСС",
    )
    add_page_break(doc)


def _s2(doc: Document) -> None:
    add_heading(doc, "2. Модель летательного аппарата FX1", 1)
    add_body(
        doc,
        "Нелинейные уравнения движения шести степеней свободы: аэродинамика, тяга, гравитация и "
        "переносные/кориолисовы ускорения модели Земли, кинематика углов и координат, динамика "
        "приводов. Состояние "
        + E(r"x \in \mathbb{R}^{25}")
        + ": скорость и угловая скорость в связанной СК, крен, курс, тангаж, дальность, высота, "
        "боковое отклонение, широта, долгота, положения приводов "
        + E(r"\delta_{T}, \delta_{V}, \delta_{N}, \delta_{E}")
        + ", ветер, вспомогательные переменные. Управление",
    )
    add_body(doc, D(r"u = \left[\delta_{T},\ \delta_{V},\ \delta_{N},\ \delta_{E}\right]^{T}"))
    add_body(
        doc,
        "Правая часть "
        + E(r"\dot{x} = f(x, u, t)")
        + " возвращает также идеальные выходы инерциальных приборов в связанной СК: абсолютную "
        "угловую скорость "
        + E(r"\omega_{bi}^{b}")
        + " и кажущееся ускорение без гравитации "
        + E(r"a_{f}^{b}")
        + " (то, что измерили бы идеальные ДУС и ДЛУ).",
    )
    add_body(
        doc,
        "Интегрирование — метод Рунге–Кутты 4-го порядка с шагом "
        + E(r"\Delta t = 10^{-3}")
        + " с:",
    )
    add_body(doc, D(r"k_{1} = f(x_{k},\ u_{k},\ t_{k})"))
    add_body(doc, D(r"k_{2} = f\left(x_{k} + \tfrac{\Delta t}{2}k_{1},\ u_{k},\ t_{k} + \tfrac{\Delta t}{2}\right)"))
    add_body(doc, D(r"k_{3} = f\left(x_{k} + \tfrac{\Delta t}{2}k_{2},\ u_{k},\ t_{k} + \tfrac{\Delta t}{2}\right)"))
    add_body(doc, D(r"k_{4} = f\left(x_{k} + \Delta t\, k_{3},\ u_{k},\ t_{k} + \Delta t\right)"))
    add_body(
        doc,
        D(r"x_{k+1} = x_{k} + \frac{\Delta t}{6}\left(k_{1} + 2k_{2} + 2k_{3} + k_{4}\right)"),
    )
    add_body(
        doc,
        "Дополнительные выходы "
        + E(r"\omega_{bi}^{b}")
        + " и "
        + E(r"a_{f}^{b}")
        + " берутся из последнего вызова правой части на шаге (конец интервала), то есть из "
        + E(r"k_{4}")
        + ".",
    )
    add_body(
        doc,
        "Автопилот удерживает высоту 500 м и скорость 80 м/с, такт 0,01 с. С "
        + E(r"t = 10")
        + " с по "
        + E(r"t = 15")
        + " с на руль высоты накладывается doublet:",
    )
    add_body(
        doc,
        D(
            r"\delta_{в}(t) = \begin{cases} +5^{\circ}, & 10 \le t < 12{,}5 \\ "
            r"-5^{\circ}, & 12{,}5 \le t < 15 \\ "
            r"\text{автопилот}, & \text{иначе} \end{cases}"
        ),
    )
    add_body(
        doc,
        "На интервале манёвра канал "
        + E(r"\delta_{V}")
        + " автопилота заморожен, чтобы не гасить вход идентификации.",
    )
    add_page_break(doc)


def _s3(doc: Document) -> None:
    add_heading(doc, "3. Модели датчиков", 1)
    add_heading(doc, "3.1. ДУС и ДЛУ", 2)
    add_body(doc, "Ошибки добавляются после FX1 и до БИНС:")
    add_body(doc, D(r"\omega_{m} = \omega_{bi}^{b} + \Delta\omega_{0} + n_{\omega}"))
    add_body(doc, D(r"a_{m} = a_{f}^{b} + \Delta a_{0} + n_{a}"))
    add_body(
        doc,
        "Смещения нуля постоянны на весь прогон, независимы по трём осям; белый шум независим на "
        "каждом шаге 0,001 с:",
    )
    add_body(doc, D(r"\Delta\omega_{0} \sim N\left(0,\ \sigma_{\omega 0}^{2}I_{3}\right), \ \ \ n_{\omega}(t_{k}) \sim N\left(0,\ \sigma_{\omega}^{2}I_{3}\right)"))
    add_body(doc, D(r"\Delta a_{0} \sim N\left(0,\ \sigma_{a0}^{2}I_{3}\right), \ \ \ n_{a}(t_{k}) \sim N\left(0,\ \sigma_{a}^{2}I_{3}\right)"))
    add_body(
        doc,
        "Для ДУС "
        + E(r"\sigma_{\omega 0} = 0{,}5\ ^{\circ}/\text{ч} = 2{,}4241\cdot 10^{-6}")
        + " рад/с и "
        + E(r"\sigma_{\omega} = 10^{-5}")
        + " рад/с; для ДЛУ "
        + E(r"\sigma_{a0} = 5\cdot 10^{-4}")
        + " м/с² и "
        + E(r"\sigma_{a} = 5\cdot 10^{-5}")
        + " м/с². Масштабные коэффициенты, перекос осей и случайные блуждания сверх этой модели не "
        "вводятся. Генераторы псевдослучайных чисел инициализируются фиксированными зёрнами "
        "(42 для инерциальных датчиков, 123 для ГНСС), поэтому прогон полностью повторяем.",
    )

    add_heading(doc, "3.2. ГНСС", 2)
    add_body(
        doc,
        "Псевдодальности и эфемериды не моделируются. Каждые "
        + E(r"\Delta t_{\mathrm{GNSS}} = 0{,}1")
        + " с к истинным "
        + E(r"\varphi, \lambda, h, V_{N}, V_{E}, V_{h}")
        + " из FX1 добавляется независимый гауссов шум:",
    )
    add_body(
        doc,
        D(
            r"np_{\mathrm{GNSS}} = \left[\varphi,\ \lambda,\ h,\ V_{N},\ V_{E},\ V_{h}\right]^{T}"
            r" + n_{\mathrm{GNSS}}, \ \ \ n_{\mathrm{GNSS}} \sim N(0,\ R)"
        ),
    )
    add_body(
        doc,
        D(
            r"R = \mathrm{diag}\left(\sigma_{\varphi}^{2},\ \sigma_{\lambda}^{2},\ \sigma_{h}^{2},"
            r"\ \sigma_{V_{N}}^{2},\ \sigma_{V_{E}}^{2},\ \sigma_{V_{h}}^{2}\right)"
        ),
    )
    add_table_caption(doc, "Таблица 2 — СКО шумов приёмника ГНСС")
    add_table(
        doc,
        ["Канал", "СКО"],
        [["φ, λ", "5·10⁻⁶ рад"], ["h", "1 м"], ["Vn, Ve, Vh", "0,2 м/с"]],
    )
    add_body(
        doc,
        "Та же матрица "
        + E(r"R")
        + " передаётся в ОФК-1 как матрица шумов измерений, то есть фильтр знает паспортную "
        "точность приёмника точно.",
    )
    add_page_break(doc)


def _s4(doc: Document) -> None:
    add_heading(doc, "4. Алгоритм механизации БИНС", 1)
    add_body(
        doc,
        "Механизация — перевод показаний ДУС и ДЛУ в скорость, координаты и ориентацию. На шаге "
        "сначала интегрируется скорость и координаты по "
        + E(r"a_{m}")
        + ", затем ориентация по "
        + E(r"\omega_{m}")
        + ". Навигационный вектор:",
    )
    add_body(doc, D(r"np = \left[V_{N},\ V_{h},\ V_{E},\ h,\ \varphi,\ \lambda\right]^{T}"))
    add_body(doc, "Ускорение в географической СК и динамика скорости:")
    add_body(doc, D(r"a^{n} = C_{b}^{n}a_{m}"))
    add_body(doc, D(r"\dot{V}^{n} = a^{n} + a_{p} + a_{k} + a_{t} + g"))
    add_body(doc, "где")
    add_body(
        doc,
        D(r"a_{p} = -\omega_{ei}^{n}\times\left(\omega_{ei}^{n}\times r^{n}\right)"),
    )
    add_body(doc, D(r"a_{k} = -2\,\omega_{ei}^{n}\times V^{n}"))
    add_body(doc, D(r"a_{t} = -\omega_{ge}^{n}\times V^{n}"))
    add_body(
        doc,
        E(r"a_{p}")
        + " — переносное ускорение от вращения Земли, "
        + E(r"a_{k}")
        + " — кориолисово, "
        + E(r"a_{t}")
        + " — от вращения географической СК, "
        + E(r"g")
        + " — гравитация earthmodel. Входящие в них векторы:",
    )
    add_body(
        doc,
        D(r"\omega_{ei}^{n} = \left[U\cos\varphi,\ U\sin\varphi,\ 0\right]^{T}, \ \ \ U = 7{,}292115\cdot 10^{-5}\ \text{рад/с}"),
    )
    add_body(
        doc,
        D(
            r"\omega_{ge}^{n} = \left[\frac{V_{E}}{R_{1}},\ \frac{V_{E}\tan\varphi}{R_{1}},"
            r"\ -\frac{V_{N}}{R_{2}}\right]^{T}"
        ),
    )
    add_body(
        doc,
        D(
            r"r^{n} = \left[-R_{1}e^{2}\sin\varphi\cos\varphi,"
            r"\ R_{1}\left(1 - e^{2}\sin^{2}\varphi\right),\ 0\right]^{T}"
        ),
    )
    add_body(
        doc,
        "Здесь "
        + E(r"R_{1}")
        + " — радиус кривизны первого вертикала, "
        + E(r"R_{2}")
        + " — меридианный радиус, оба с учётом высоты; "
        + E(r"e^{2}")
        + " — квадрат эксцентриситета эллипсоида Красовского. Гравитация "
        + E(r"g")
        + " вычисляется по монопольной и квадрупольной составляющим поля с переходом от "
        "геодезической широты к геоцентрической.",
    )
    add_body(doc, "Кинематика:")
    add_body(
        doc,
        D(
            r"\dot{h} = V_{h}, \ \ \ \dot{\varphi} = \frac{V_{N}}{R_{2}}, \ \ \ "
            r"\dot{\lambda} = \frac{V_{E}}{R_{1}\cos\varphi}"
        ),
    )
    add_body(
        doc,
        "Дискретизация навигационного канала — явный метод Эйлера с шагом "
        + E(r"\Delta t = 10^{-3}")
        + " с: "
        + E(r"np_{k+1} = np_{k} + \dot{np}_{k}\,\Delta t")
        + ".",
    )
    add_body(
        doc,
        "Ориентация. Угловая скорость географической СК относительно инерциальной переводится в "
        "связанную, относительная угловая скорость прибора:",
    )
    add_body(doc, D(r"\omega_{ni}^{n} = \omega_{ne}^{n} + \omega_{ei}^{n}, \ \ \ \omega_{ne}^{n} = \omega_{ge}^{n}"))
    add_body(doc, D(r"\omega_{bn}^{b} = \omega_{m} - \left(C_{b}^{n}\right)^{T}\omega_{ni}^{n}"))
    add_body(
        doc,
        "Кватернион интегрируется с нормировочной поправкой "
        + E(r"\left(1 - \|q\|^{2}\right)")
        + ":",
    )
    add_body(
        doc,
        D(
            r"\dot{q} = \tfrac{1}{2}\,q \circ \omega_{bn}^{b}"
            r" + \tfrac{1}{2}\left(1 - \|q\|^{2}\right)q"
        ),
    )
    add_body(doc, "или покомпонентно, при " + E(r"\omega_{bn}^{b} = [\omega_{x},\ \omega_{y},\ \omega_{z}]^{T}") + ":")
    add_body(
        doc,
        D(
            r"\dot{q}_{0} = \tfrac{1}{2}\left(-q_{1}\omega_{x} - q_{2}\omega_{y} - q_{3}\omega_{z}"
            r" + q_{0}\left(1 - \|q\|^{2}\right)\right)"
        ),
    )
    add_body(
        doc,
        D(
            r"\dot{q}_{1} = \tfrac{1}{2}\left(q_{0}\omega_{x} - q_{3}\omega_{y} + q_{2}\omega_{z}"
            r" + q_{1}\left(1 - \|q\|^{2}\right)\right)"
        ),
    )
    add_body(
        doc,
        D(
            r"\dot{q}_{2} = \tfrac{1}{2}\left(q_{3}\omega_{x} + q_{0}\omega_{y} - q_{1}\omega_{z}"
            r" + q_{2}\left(1 - \|q\|^{2}\right)\right)"
        ),
    )
    add_body(
        doc,
        D(
            r"\dot{q}_{3} = \tfrac{1}{2}\left(-q_{2}\omega_{x} + q_{1}\omega_{y} + q_{0}\omega_{z}"
            r" + q_{3}\left(1 - \|q\|^{2}\right)\right)"
        ),
    )
    add_body(doc, "Матрица " + E(r"C_{b}^{n}") + " восстанавливается из кватерниона:")
    add_body(
        doc,
        D(
            r"C_{n}^{b} = \begin{pmatrix}"
            r"1 - 2(q_{2}^{2}+q_{3}^{2}) & 2(q_{1}q_{2}+q_{3}q_{0}) & 2(q_{1}q_{3}-q_{2}q_{0}) \\"
            r"2(q_{1}q_{2}-q_{3}q_{0}) & 1 - 2(q_{1}^{2}+q_{3}^{2}) & 2(q_{2}q_{3}+q_{1}q_{0}) \\"
            r"2(q_{1}q_{3}+q_{2}q_{0}) & 2(q_{2}q_{3}-q_{1}q_{0}) & 1 - 2(q_{1}^{2}+q_{2}^{2})"
            r"\end{pmatrix}, \ \ \ C_{b}^{n} = \left(C_{n}^{b}\right)^{T}"
        ),
    )
    add_body(
        doc,
        "ОФК-1 "
        + E(r"q")
        + " и "
        + E(r"C_{b}^{n}")
        + " не корректирует. Инициализация: "
        + E(r"q")
        + ", "
        + E(r"C_{b}^{n}")
        + " и "
        + E(r"np")
        + " совпадают с истиной FX1 (без ошибок датчиков), далее расхождение накапливается из "
        "ошибок ДУС и ДЛУ.",
    )
    add_page_break(doc)


def _s5(doc: Document) -> None:
    add_heading(doc, "5. Алгоритм ОФК-1 (навигация)", 1)
    add_body(
        doc,
        "ОФК-1 оценивает ошибки решения БИНС. Каждые 0,1 с фильтр сравнивает БИНС с ГНСС, "
        "оценивает эти ошибки и вычитает их из навигационного вектора. Модель самолёта FX1 при "
        "этом не меняется. Пятнадцать ошибок навигации. Первые тринадцать — ошибки ориентации, "
        "северная и восточная скорости, широта и долгота, смещения нуля ДУС и ДЛУ. Ещё две "
        "компоненты — ошибка вертикальной скорости и ошибка высоты. Они добавлены, чтобы по ГНСС "
        "можно было править вертикальный канал:",
    )
    add_body(
        doc,
        D(
            r"X = \left[\alpha_{x},\alpha_{y},\alpha_{z},\ \delta V_{N},\delta V_{E},"
            r"\ \delta\varphi,\delta\lambda,\ \Delta a_{0x},\Delta a_{0y},\Delta a_{0z},"
            r"\ \Delta\omega_{0x},\Delta\omega_{0y},\Delta\omega_{0z},"
            r"\ \delta V_{h},\ \delta h\right]^{T}"
        ),
    )
    add_body(doc, "Инновация:")
    add_body(
        doc,
        D(
            r"z = \left[\varphi_{Б}-\varphi_{Г},\ \lambda_{Б}-\lambda_{Г},"
            r"\ V_{N,Б}-V_{N,Г},\ V_{E,Б}-V_{E,Г},"
            r"\ V_{h,Б}-V_{h,Г},\ h_{Б}-h_{Г}\right]^{T}"
        ),
    )
    add_body(
        doc,
        "Это невязка "
        + E(r"z")
        + ". Если решения совпали, "
        + E(r"z")
        + " близка к нулю (остаётся шум ГНСС). Между тактами ГНСС ошибки развиваются по линейной "
        "модели: производная состояния равна "
        + E(r"F\!\cdot\!X")
        + " плюс шум датчиков через "
        + E(r"G")
        + ". Невязка "
        + E(r"z")
        + " связана с ошибками матрицей "
        + E(r"H")
        + " плюс шум ГНСС. Матрицы "
        + E(r"F, G, H, Q")
        + " каждый такт считаются заново из текущего ускорения ДЛУ, ориентации и вектора БИНС.",
    )
    add_body(doc, "Непрерывная модель ошибок на такте ГНСС:")
    add_body(doc, D(r"\dot{X} = F X + G w, \ \ \ z = H X + v"))
    add_body(
        doc,
        D(r"w \sim N(0,\ Q), \ \ \ v \sim N(0,\ R)"),
    )
    add_body(
        doc,
        "Матрица "
        + E(r"H")
        + " для исходных тринадцати состояний «видит» только ошибки широты, долготы, "
        + E(r"V_{N}")
        + " и "
        + E(r"V_{E}")
        + ". Для двух добавленных состояний: производная ошибки высоты равна ошибке вертикальной "
        "скорости, "
        + E(r"F_{14,13} = 1")
        + ". Вертикальный канал в "
        + E(r"F")
        + " не связан с горизонтальными скоростями — иначе через ковариацию портилась "
        "горизонтальная оценка; связь оставлена только с ошибками ориентации и смещением ДЛУ:",
    )
    add_body(
        doc,
        D(r"F_{13,\,0:3} = \left[n_{y},\ -n_{x},\ 0\right], \ \ \ n = C_{2}\,a_{m}"),
    )
    add_body(doc, "Ненулевые элементы матрицы наблюдения:")
    add_body(
        doc,
        D(
            r"H_{1,6} = H_{2,7} = H_{3,4} = H_{4,5} = H_{5,14} = H_{6,15} = 1"
        ),
    )
    add_body(doc, "Процессный шум:")
    add_body(
        doc,
        D(
            r"Q = \frac{1}{\Delta t_{\mathrm{GNSS}}}\,\mathrm{diag}"
            r"\left(w_{\omega}^{2},w_{\omega}^{2},w_{\omega}^{2},"
            r"\ w_{a}^{2},w_{a}^{2},w_{a}^{2}\right)"
        ),
    )
    add_body(
        doc,
        "где "
        + E(r"w_{\omega} = 2\cdot 10^{-9}")
        + " и "
        + E(r"w_{a} = 2{,}5\cdot 10^{-6}")
        + ". Матрица "
        + E(r"R")
        + " собирается из паспортных СКО ГНСС (таблица 2) с нижним порогом по каждому каналу, "
        "переставленных под порядок строк "
        + E(r"H")
        + ".",
    )
    add_body(doc, "Предсказание, " + E(r"\Delta t = \Delta t_{\mathrm{GNSS}}") + ":")
    add_body(
        doc,
        D(r"\Phi = I + F\Delta t + \tfrac{1}{2}\left(F\Delta t\right)^{2}, \ \ \ \Gamma = G\Delta t"),
    )
    add_body(doc, D(r"\hat{X}^{-} = \Phi\hat{X}^{+}"))
    add_body(doc, D(r"P^{-} = \Phi P^{+}\Phi^{T} + \Gamma Q \Gamma^{T}"))
    add_body(doc, "Обновление:")
    add_body(doc, D(r"S = H P^{-}H^{T} + R"))
    add_body(doc, D(r"K = P^{-}H^{T}S^{-1}"))
    add_body(doc, D(r"\hat{X}^{+} = \hat{X}^{-} + K\left(z - H\hat{X}^{-}\right)"))
    add_body(
        doc,
        D(r"P^{+} = \left(I - KH\right)P^{-}\left(I - KH\right)^{T} + K R K^{T}"),
    )
    add_body(
        doc,
        "Ковариация обновляется в форме Джозефа — она сохраняет симметрию и положительную "
        "определённость при накоплении ошибок округления.",
    )
    add_body(
        doc,
        "Оценённые ошибки вычитаются только из решения БИНС: "
        + E(r"V_{N}, V_{E}, \varphi, \lambda, V_{h}, h")
        + ". Состояние FX1 не трогается. Кватернион и матрица ориентации этим шагом не "
        "корректируются:",
    )
    add_body(
        doc,
        D(
            r"V_{N} \leftarrow V_{N} - \hat{X}_{4}, \ \ \ V_{E} \leftarrow V_{E} - \hat{X}_{5}, "
            r"\ \ \ \varphi \leftarrow \varphi - \hat{X}_{6}, \ \ \ "
            r"\lambda \leftarrow \lambda - \hat{X}_{7}, \ \ \ "
            r"V_{h} \leftarrow V_{h} - \hat{X}_{14}, \ \ \ h \leftarrow h - \hat{X}_{15}"
        ),
    )
    add_body(
        doc,
        "После вычитания вектор оценки ошибок обнуляется: "
        + E(r"\hat{X} := 0")
        + " — ошибка уже «вынута» из БИНС, дальше копятся только новые. Ковариация "
        + E(r"P")
        + " при этом не сбрасывается и служит расчётной оценкой достигнутой точности. Начальная "
        "ковариация:",
    )
    add_body(
        doc,
        D(
            r"P_{0} = \mathrm{diag}\left(10^{3}I_{4},\ 10^{-2}I_{6},\ 10^{-4}I_{5}\right),"
            r"\ \ \ P_{0,14} = 0{,}04, \ \ \ P_{0,15} = 1{,}0"
        ),
    )
    add_body(
        doc,
        "Поправки для вертикального канала согласованы с шумом ГНСС: "
        + E(r"\sigma_{V_{h}} = 0{,}2")
        + " м/с и "
        + E(r"\sigma_{h} = 1")
        + " м. На первом такте ГНСС скорости и высота БИНС сначала копируются из ГНСС, затем "
        "выполняется фильтр.",
    )
    add_page_break(doc)


def _s6(doc: Document) -> None:
    add_heading(doc, "6. Модель и алгоритм ОФК-2 (идентификация)", 1)
    add_heading(doc, "6.1. Модель короткопериодического движения", 2)
    add_body(doc, "Линейная модель в приращениях от точки балансировки " + E(r"t = 0") + ":")
    add_body(
        doc,
        D(
            r"\dot{\alpha} = L_{\alpha}\delta\alpha + L_{q}q + L_{\delta e}\delta\delta_{e}"
            r" + L_{V}\delta V + L_{\theta}\delta\theta"
        ),
    )
    add_body(
        doc,
        D(
            r"\dot{q} = M_{\alpha}\delta\alpha + M_{q}q + M_{\delta e}\delta\delta_{e}"
            r" + M_{V}\delta V + M_{\theta}\delta\theta"
        ),
    )
    add_body(
        doc,
        D(
            r"a_{z} = a_{z0} - \frac{V}{g}\left[L_{\alpha}\delta\alpha + \left(L_{q}-1\right)q"
            r" + L_{\delta e}\delta\delta_{e} + L_{V}\delta V + L_{\theta}\delta\theta\right]"
        ),
    )
    add_body(
        doc,
        "Член "
        + E(r"\left(L_{q}-1\right)q")
        + " отделяет кинематику "
        + E(r"\dot{\alpha} \approx q")
        + " от аэродинамического вклада в нормальную перегрузку. На балансировке якобиан FX1 даёт "
        + E(r"L_{q} = 1")
        + ".",
    )
    add_body(doc, "Вектор состояния фильтра (оценка десяти коэффициентов):")
    add_body(
        doc,
        D(
            r"\hat{\theta} = \left[L_{\alpha},L_{q},L_{\delta e},L_{V},L_{\theta},"
            r"\ M_{\alpha},M_{q},M_{\delta e},M_{V},M_{\theta}\right]^{T}"
        ),
    )
    add_body(doc, "Регрессор:")
    add_body(
        doc,
        D(
            r"\phi = \left[\delta\alpha,\ q,\ \delta\delta_{e},\ \delta V,\ \delta\theta\right]^{T}"
            r" = \left[\alpha-\alpha_{0},\ q,\ \delta_{e}-\delta_{e0},"
            r"\ V-V_{0},\ \theta-\theta_{0}\right]^{T}"
        ),
    )
    add_body(
        doc,
        "Источники сигналов. Скорость в связанной СК "
        + E(r"v^{b} = \left(C_{b}^{n}\right)^{T}\left[V_{N},\ V_{h},\ V_{E}\right]^{T}")
        + ",",
    )
    add_body(
        doc,
        D(
            r"\alpha = -\arctan_{2}\left(v_{y}^{b},\ v_{x}^{b}\right), \ \ \ "
            r"V = \left\|v^{b}\right\|"
        ),
    )
    add_body(
        doc,
        "Тангаж из "
        + E(r"C_{b}^{n}")
        + ". "
        + E(r"q = \omega_{m,z}")
        + ", "
        + E(r"a_{z} = a_{m,y}/g")
        + ", "
        + E(r"g = 9{,}80665")
        + " м/с². Руль "
        + E(r"\delta_{e}")
        + " — с положения привода FX1, перевод из градусов в радианы. Величина "
        + E(r"a_{z0}")
        + " равна первому измеренному "
        + E(r"a_{z}")
        + " и далее не обновляется.",
    )

    add_heading(doc, "6.2. Эталон (численный якобиан FX1)", 2)
    add_body(
        doc,
        "Эталон — центральные разности правых частей FX1. При фиксированных остальных аргументах:",
    )
    add_body(
        doc,
        D(
            r"L_{\alpha} = \frac{\dot{\alpha}\left(\alpha_{0}+\Delta\alpha\right)"
            r" - \dot{\alpha}\left(\alpha_{0}-\Delta\alpha\right)}{2\Delta\alpha}, \ \ \ "
            r"M_{\alpha} = \frac{\dot{q}\left(\alpha_{0}+\Delta\alpha\right)"
            r" - \dot{q}\left(\alpha_{0}-\Delta\alpha\right)}{2\Delta\alpha}"
        ),
    )
    add_body(
        doc,
        "аналогично "
        + E(r"L_{q}, L_{\delta e}, L_{V}, L_{\theta}")
        + " и все "
        + E(r"M_{*}")
        + " по "
        + E(r"\dot{q}")
        + ". Шаги: "
        + E(r"\Delta\alpha = \Delta q = \Delta\delta_{e} = \Delta\theta = 10^{-4}")
        + ", "
        + E(r"\Delta V = 0{,}05")
        + " м/с. "
        + E(r"\theta_{\mathrm{trim}}")
        + " считается один раз в "
        + E(r"(x_{0}, u_{0})")
        + ". "
        + E(r"\theta_{\mathrm{FX1}}(t)")
        + " пересчитывается каждый такт ОФК-2 из текущего "
        + E(r"(x, u)")
        + ". Фильтр эталон не использует.",
    )
    add_table_caption(doc, "Таблица 3 — Точка балансировки, эталон и старт фильтра")
    add_table(
        doc,
        ["Параметр", "θ_trim", "Старт фильтра 1,3·θ_trim"],
        [
            ["α₀, θ₀", "0,212 рад", "—"],
            ["q₀", "0", "—"],
            ["δe₀", "−0,125 рад", "—"],
            ["V₀", "80 м/с", "—"],
            ["Lα", "−0,4961", "−0,6449"],
            ["Lq", "1,000", "1,300"],
            ["Lδe", "−0,03118", "−0,04054"],
            ["Lv", "−0,002958", "−0,003845"],
            ["Lθ", "≈0", "≈0"],
            ["Mα", "−0,2791", "−0,3628"],
            ["Mq", "−0,6877", "−0,8940"],
            ["Mδe", "−0,4244", "−0,5517"],
            ["Mv", "0,000570", "0,000741"],
            ["Mθ", "≈0", "≈0"],
        ],
    )

    add_heading(doc, "6.3. Equation-error EKF", 2)
    add_body(
        doc,
        "Как из "
        + E(r"\alpha, q, \delta_{e}, V, \theta, a_{z}")
        + " получаются 10 коэффициентов. Коэффициенты считаются постоянными, поэтому",
    )
    add_body(doc, D(r"\dot{\theta} = 0, \ \ Q = 0 \quad\Rightarrow\ \ \hat{\theta}^{-} = \hat{\theta}^{+}, \ \ P^{-} = P^{+}"))
    add_body(
        doc,
        E(r"\dot{\alpha}")
        + " и "
        + E(r"\dot{q}")
        + " не измеряются напрямую — берутся конечными разностями за "
        + E(r"\Delta t_{2} = 0{,}02")
        + " с:",
    )
    add_body(
        doc,
        D(
            r"\dot{\alpha}_{fd} = \frac{\delta\alpha_{k} - \delta\alpha_{k-1}}{\Delta t_{2}},"
            r"\ \ \ \dot{q}_{fd} = \frac{q_{k} - q_{k-1}}{\Delta t_{2}}"
        ),
    )
    add_body(doc, "Измерение " + E(r"Z = \left[\dot{\alpha}_{fd},\ \dot{q}_{fd},\ a_{z}\right]^{T}") + ". Матрица наблюдения " + E(r"H \in \mathbb{R}^{3\times 10}") + ":")
    add_body(
        doc,
        D(
            r"H = \begin{pmatrix}"
            r"\phi^{T} & 0_{1\times 5} \\"
            r"0_{1\times 5} & \phi^{T} \\"
            r"-\dfrac{V}{g}\phi^{T} & 0_{1\times 5}"
            r"\end{pmatrix}"
        ),
    )
    add_body(
        doc,
        "Первая строка — уравнение "
        + E(r"\dot{\alpha}")
        + " (коэффициенты "
        + E(r"L")
        + "), вторая — "
        + E(r"\dot{q}")
        + " (коэффициенты "
        + E(r"M")
        + "), третья — "
        + E(r"a_{z}")
        + " только через "
        + E(r"L")
        + ". Прогноз измерения:",
    )
    add_body(
        doc,
        D(
            r"\hat{Z} = \begin{pmatrix} L^{T}\phi \\ M^{T}\phi \\"
            r"a_{z0} - \dfrac{V}{g}\left[L_{\alpha}\delta\alpha + \left(L_{q}-1\right)q"
            r" + L_{\delta e}\delta\delta_{e} + L_{V}\delta V + L_{\theta}\delta\theta\right]"
            r"\end{pmatrix}, \ \ \ \delta = Z - \hat{Z}"
        ),
    )
    add_body(
        doc,
        "Единица в третьей строке прогноза — известное слагаемое, в "
        + E(r"H")
        + " она не входит, поэтому производная по "
        + E(r"L_{q}")
        + " равна "
        + E(r"-(V/g)\,q")
        + ".",
    )
    add_body(doc, D(r"R = \mathrm{diag}\left(0{,}025^{2},\ 0{,}002^{2},\ 0{,}023^{2}\right)"))
    add_body(
        doc,
        "Начальная "
        + E(r"P_{0}")
        + " диагональная: "
        + E(r"3\sigma")
        + " покрывает смещение 30 % "
        + E(r"\left|\theta_{\mathrm{trim}}\right|")
        + " с абсолютным полом по каналам:",
    )
    add_body(
        doc,
        D(
            r"\sigma_{i} = \max\left(\frac{0{,}3\left|\theta_{0,i}\right|}{3},"
            r"\ 0{,}2\,c_{i}\right), \ \ \ P_{0} = \mathrm{diag}\left(\sigma_{i}^{2}\right)"
        ),
    )
    add_body(
        doc,
        "где "
        + E(r"c = [0{,}05;\ 0{,}05;\ 0{,}02;\ 0{,}01;\ 0{,}01;\ 0{,}05;\ 0{,}05;\ 0{,}05;\ 0{,}01;\ 0{,}01]")
        + " — абсолютный пол по каналам.",
    )
    add_body(
        doc,
        "Алгоритм шага. Предсказание: "
        + E(r"\hat{\theta}^{-} = \hat{\theta}^{+}")
        + ", "
        + E(r"P^{-} = P^{+}")
        + " (Q = 0). Если нет предыдущего такта или нет возбуждения, обновление не выполняется. "
        "Возбуждение, если выполнено хотя бы одно условие (таблица 4). Иначе шаг пропускается. "
        "При обновлении — стандартный EKF по "
        + E(r"Z, H, R")
        + ", как у ОФК-1:",
    )
    add_body(doc, D(r"K = P^{-}H^{T}\left(H P^{-}H^{T} + R\right)^{-1}"))
    add_body(doc, D(r"\hat{\theta}^{+} = \hat{\theta}^{-} + K\delta"))
    add_body(
        doc,
        D(r"P^{+} = \left(I - KH\right)P^{-}\left(I - KH\right)^{T} + K R K^{T}"),
    )
    add_table_caption(doc, "Таблица 4 — Пороги возбуждения для обновления ОФК-2")
    add_table(
        doc,
        ["Регрессор", "Порог", "Единицы"],
        [
            ["|δδe|", "0,003", "рад"],
            ["|q|", "0,005", "рад/с"],
            ["|δα|", "0,008", "рад"],
            ["|δV|", "0,05", "м/с"],
            ["|δθ|", "0,003", "рад"],
        ],
    )
    add_body(
        doc,
        "Метрики: ошибка относительно балансировки "
        + E(r"\Delta_{\mathrm{trim}} = \hat{\theta} - \theta_{\mathrm{trim}}")
        + ", ошибка относительно текущего эталона "
        + E(r"\Delta_{\mathrm{inst}} = \hat{\theta} - \theta_{\mathrm{FX1}}(t)")
        + ". "
        + E(r"L_{\theta}")
        + " и "
        + E(r"M_{\theta}")
        + " остаются в состоянии фильтра, на балансировке "
        + E(r"\approx 0")
        + "; в отчётных графиках: "
        + E(r"L_{\alpha}, L_{q}, L_{\delta e}, L_{V}, M_{\alpha}, M_{q}, M_{\delta e}, M_{V}")
        + ".",
    )
    add_page_break(doc)


def _s7(doc: Document, nav: dict, coeff: list) -> None:
    add_heading(doc, "7. Результаты прогона", 1)
    add_body(
        doc,
        "Ниже приведены результаты прогона со сценарием из разделов 1–3: "
        + E(r"t = 60")
        + " с, "
        + E(r"\Delta t = 10^{-3}")
        + " с, doublet "
        + E(r"\pm 5^{\circ}")
        + " на интервале 10…15 с, зёрна генераторов фиксированы.",
    )

    add_heading(doc, "7.1. Траектория и навигационное решение", 2)
    add_figure(doc, PLOTS / "latitude_vs_longitude.png", 12.0)
    add_caption(doc, "Рисунок 3 — Траектория полёта по истине FX1")
    add_body(
        doc,
        "На рисунке 4 совмещены истина FX1, счисление БИНС до коррекции, решение после ОФК-1 и "
        "измерения ГНСС. Решение после ОФК-1 ложится на истину в пределах толщины линии, "
        "измерения ГНСС образуют вокруг неё шумовое облако.",
    )
    add_figure(doc, PLOTS / "bins_gnss_full.png", 16.5)
    add_caption(
        doc,
        "Рисунок 4 — Навигационные параметры: истина, БИНС, решение после ОФК-1 и ГНСС",
    )
    add_page_break(doc)

    add_heading(doc, "7.2. Ошибки ОФК-1 и полоса ±3σ", 2)
    add_body(
        doc,
        "Полоса строится по диагонали апостериорной ковариации фильтра: "
        + E(r"\pm 3\sqrt{P_{ii}}")
        + " для каналов "
        + E(r"\delta\varphi, \delta\lambda, \delta h, \delta V_{N}, \delta V_{E}, \delta V_{h}")
        + ". Ошибка считается относительно истины FX1: "
        + E(r"\delta = x_{\mathrm{ОФК-1}} - x_{\mathrm{истина}}")
        + ".",
    )
    add_figure(doc, PLOTS / "ofk_error_vs_three_sigma_P.png", 15.5)
    add_caption(doc, "Рисунок 5 — Ошибки навигации после ОФК-1 и полосы ±3√Pᵢᵢ")
    add_body(
        doc,
        "Переходный процесс занимает около 5 с по скоростям и около 10 с по координатам: "
        "завышенная начальная ковариация сходится к установившемуся уровню. Показатели на "
        "установившемся участке приведены в таблице 5; ошибки по широте и долготе пересчитаны в "
        "метры по радиусу Земли.",
    )
    order = ["fi", "lam", "h", "vn", "ve", "vh"]
    rows = []
    for key in order:
        s = nav[key]
        rows.append(
            [
                f"{s['name']}, {s['unit']}",
                _fmt(s["rms_ofk"]),
                _fmt(s["max_ofk"]),
                _fmt(s["sigma3_end"]),
                _pct(s["inside"]),
            ]
        )
    add_table_caption(
        doc, "Таблица 5 — Ошибки навигации после ОФК-1 на интервале t = 5…60 с"
    )
    add_table(
        doc,
        ["Канал", "СКО", "Максимум модуля", "3σ в конце прогона", "Внутри ±3σ, %"],
        rows,
    )
    add_body(
        doc,
        "Координаты держатся на уровне единиц метров, скорости — на уровне сотых долей метра в "
        "секунду, то есть заметно точнее одиночного измерения ГНСС (≈32 м по координатам и "
        "0,2 м/с по скорости). Все шесть каналов лежат внутри расчётной полосы, значит "
        + E(r"Q")
        + " и "
        + E(r"R")
        + " согласованы с фактическими ошибками. Отдельно отмечу: обратная связь вводится каждые "
        "0,1 с, поэтому «БИНС до коррекции» на графиках — это дрейф за один такт, а не автономный "
        "уход БИНС; для оценки автономного режима нужен отдельный прогон с отключённой обратной "
        "связью.",
    )
    add_page_break(doc)

    add_heading(doc, "7.3. Идентификация коэффициентов", 2)
    add_body(
        doc,
        "Старт фильтра смещён на 30 % от "
        + E(r"\theta_{\mathrm{trim}}")
        + " (таблица 3). Основное уточнение происходит на интервале манёвра 10…15 с, когда "
        "выполняются условия возбуждения из таблицы 4; вне манёвра обновление пропускается и "
        "оценки держатся на месте.",
    )
    add_figure(doc, PLOTS / "ofk2_params_time.png", 15.0)
    add_caption(doc, "Рисунок 6 — Сходимость оценок L*, M* к эталону")
    add_figure(doc, PLOTS / "ofk2_params_three_sigma.png", 15.0)
    add_caption(doc, "Рисунок 7 — Оценки L*, M* и полосы ±3σ")
    add_figure(doc, PLOTS / "ofk2_ekf.png", 15.0)
    add_caption(
        doc,
        "Рисунок 8 — Входы ОФК-2 (α, q, az), невязки equation-error и Δ = оценка − теория",
    )
    rows = []
    for c in coeff:
        rows.append(
            [
                c["name"],
                _fmt(c["theory"], 4),
                _fmt(c["start"], 4),
                _fmt(c["est"], 4),
                _fmt(c["delta"], 4),
                _fmt(c["sigma3"], 4),
                _pct(c["rel"]),
            ]
        )
    add_table_caption(doc, "Таблица 6 — Результат идентификации на t = 60 с")
    add_table(
        doc,
        ["Коэффициент", "Эталон", "Старт", "Оценка", "Δ", "3σ", "|Δ|/|эталон|, %"],
        rows,
    )
    add_body(
        doc,
        "Лучше всего восстанавливаются коэффициенты, которые напрямую возбуждает руль высоты: "
        + E(r"L_{q}, M_{\delta e}, M_{q}, L_{\alpha}")
        + " — отклонение от эталона в пределах нескольких процентов. Хуже всего "
        + E(r"M_{V}")
        + " и "
        + E(r"L_{\delta e}")
        + ": они малы по модулю и слабо наблюдаемы на doublet, поэтому относительная ошибка "
        "велика при малой абсолютной. Для большинства каналов "
        + E(r"\left|\Delta\right|")
        + " сопоставима с расчётной "
        + E(r"3\sigma")
        + ", то есть фильтр не переоценивает свою точность.",
    )


def _s8(doc: Document, nav: dict, coeff: list) -> None:
    add_heading(doc, "8. Итог и что осталось", 1)
    max_pos = max(nav["fi"]["rms_ofk"], nav["lam"]["rms_ofk"])
    max_vel = max(nav["vn"]["rms_ofk"], nav["ve"]["rms_ofk"], nav["vh"]["rms_ofk"])
    worst = max(coeff, key=lambda c: c["rel"])
    best = min(coeff, key=lambda c: c["rel"])
    add_body(
        doc,
        "Контур собран целиком и считает за один прогон: FX1, ошибки ДУС/ДЛУ, механизацию БИНС, "
        "ГНСС, ОФК-1 и ОФК-2 на четырёх разных тактах.",
    )
    add_body(
        doc,
        "ОФК-1 даёт координаты с СКО не хуже "
        + _fmt(max_pos, 2)
        + " м и скорости не хуже "
        + _fmt(max_vel, 3)
        + " м/с; все каналы внутри "
        + E(r"\pm 3\sqrt{P_{ii}}")
        + ".",
    )
    add_body(
        doc,
        "ОФК-2 при старте с 30 % смещением сводит коэффициенты к эталону с отклонением от "
        + _pct(best["rel"])
        + " % ("
        + best["name"]
        + ") до "
        + _pct(worst["rel"])
        + " % ("
        + worst["name"]
        + "); большие относительные отклонения приходятся на слабонаблюдаемые каналы.",
    )
    add_body(
        doc,
        "Что не сделано и стоит добавить: прогон с отключённой обратной связью ОФК-1 для оценки "
        "автономного дрейфа БИНС; ошибки начальной выставки БИНС; пропадания сигнала ГНСС; "
        "сравнение информативности разных манёвров (doublet против 3-2-1-1) для идентификации; "
        "масштабные коэффициенты и перекос осей в модели датчиков.",
    )


def build() -> Path:
    if not STATS.exists():
        raise SystemExit(f"Нет {STATS}. Сначала: python docs/_report_stats.py")
    st = json.loads(STATS.read_text(encoding="utf-8"))
    nav, coeff = st["nav"], st["coeff"]

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, align="center", space_after=18)
    run = p.add_run("Модели и алгоритмы численного моделирования БНК и идентификации")
    set_run_font(run, size=16, bold=True)

    _s1(doc)
    _s2(doc)
    _s3(doc)
    _s4(doc)
    _s5(doc)
    _s6(doc)
    _s7(doc, nav, coeff)
    _s8(doc, nav, coeff)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
