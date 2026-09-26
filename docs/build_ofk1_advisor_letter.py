"""Письмо научруку: модели контура до ОФК-1. Запуск: poetry run python docs/build_ofk1_advisor_letter.py"""

from __future__ import annotations

import re
from pathlib import Path

import latex2mathml.converter
import mathml2omml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
OUT = Path(__file__).resolve().parent / "ОФК-1_модели_для_научника.docx"


def latex_to_omath_xml(latex: str) -> str:
    mml = latex2mathml.converter.convert(latex)
    omml = mathml2omml.convert(mml)
    if "xmlns:m=" not in omml:
        omml = omml.replace("<m:oMath>", f'<m:oMath xmlns:m="{M_NS}">', 1)
    return omml


def add_omath(paragraph, latex: str) -> None:
    omml = latex_to_omath_xml(latex)
    omml = re.sub(r'\sxmlns:m="[^"]*"', "", omml)
    wrapped = (
        f'<m:oMath xmlns:m="{M_NS}">'
        + omml.replace("<m:oMath>", "").replace("</m:oMath>", "")
        + "</m:oMath>"
    )
    paragraph._p.append(parse_xml(wrapped))


def add_display(doc: Document, latex: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    add_omath(p, latex)


def add_runs_with_math(p, text: str) -> None:
    parts = re.split(r"(\$[^$]+\$)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("$") and part.endswith("$") and len(part) > 2:
            add_omath(p, part[1:-1])
        else:
            p.add_run(part)


def set_run_font(run, *, bold=False, size=12, name="Times New Roman"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    r = run._element
    rpr = r.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)


def heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    add_runs_with_math(p, text)
    for run in p.runs:
        set_run_font(run, bold=True, size=14 if level == 1 else 13)


def para(doc: Document, text: str, *, first_indent: bool = False) -> None:
    p = doc.add_paragraph()
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    add_runs_with_math(p, text)
    for run in p.runs:
        set_run_font(run)


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        add_runs_with_math(p, item)
        for run in p.runs:
            set_run_font(run)


def build() -> Path:
    doc = Document()
    style_doc(doc)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Модели контура до ОФК-1")
    set_run_font(r, bold=True, size=16)

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run("самолёт → датчики → БИНС и ГНСС → разность → ОФК-1")
    set_run_font(r, size=12)
    r.italic = True

    para(doc, "Добрый день.")
    para(
        doc,
        "Прошу проверить модели до ОФК-1. ОФК-2 здесь нет. Ниже — как устроено в программе: "
        "что куда идёт и какие уравнения считаются.",
        first_indent=True,
    )

    heading(doc, "0. Как связаны блоки")
    para(
        doc,
        "Модель самолёта (FX1) считает полёт каждый 1 мс и отдаёт два разных набора чисел.",
    )
    bullets(
        doc,
        [
            "В датчики, а потом в БИНС: идеальная угловая скорость $\\boldsymbol{\\omega}$ и идеальное "
            "удельное ускорение $\\mathbf{a}$ в осях самолёта.",
            "В ГНСС: истинные широта, долгота, высота и три скорости "
            "$\\varphi$, $\\lambda$, $h$, $V_N$, $V_E$, $V_h$.",
            "На идеал $\\boldsymbol{\\omega}$, $\\mathbf{a}$ навешиваются ошибки ДУС и ДЛУ. "
            "В БИНС приходят уже сырые показания $\\boldsymbol{\\omega}_m$, $\\mathbf{a}_m$.",
            "БИНС считает сам, 1000 раз в секунду. Оценки ОФК-1 в БИНС не подставляем.",
            "БИНС выдаёт навигационный вектор $Y_{\\mathrm{БИНС}}$, кватернион $q$ и матрицу ориентации $C_b^n$.",
            "ГНСС работает 10 раз в секунду: к истине FX1 добавляется белый шум.",
            "ОФК-1 видит разность $\\varepsilon = Y_{\\mathrm{БИНС}} - Y_{\\mathrm{ГНСС}}$, "
            "а также $C_b^n$ и $\\mathbf{a}_m$. На выходе — 15 оценок ошибок и матрица $P$.",
        ],
    )

    heading(doc, "1. Самолёт (FX1)")
    para(
        doc,
        "Это полная нелинейная модель самолёта. Шаг 1 мс. Для БИНС и ГНСС нужны не все "
        "внутренние переменные, а только то, что уходит наружу.",
    )
    heading(doc, "1.1. Что уходит в датчики", 2)
    para(
        doc,
        "$\\boldsymbol{\\omega}=[\\omega_x,\\omega_y,\\omega_z]$ — как самолёт крутится относительно "
        "инерциального пространства (из состояния модели).",
    )
    para(
        doc,
        "$\\mathbf{a}=[a_x,a_y,a_z]$ — удельное ускорение: сила тяги и аэродинамики, делённая на массу. "
        "Силы тяжести в этом векторе нет. Тяжесть добавит уже БИНС. Если вписать $\\mathbf{g}$ сюда, "
        "БИНС вычтет её второй раз, и вертикаль поедет.",
    )
    heading(doc, "1.2. Что уходит в ГНСС", 2)
    para(
        doc,
        "Широта, долгота и высота $\\varphi$, $\\lambda$, $h$ берутся прямо из состояния FX1. "
        "Скорость в модели записана в осях самолёта $V_x$, $V_y$, $V_z$. Для ГНСС её надо "
        "перевести в географические оси (север, вверх, восток) матрицей углов самолёта "
        "$C_g^b(\\vartheta,\\gamma,\\psi)$:",
    )
    add_display(
        doc,
        r"\begin{bmatrix} V_N \\ V_h \\ V_E \end{bmatrix}"
        r"=(C_g^b)^{\top}"
        r"\begin{bmatrix} V_x \\ V_y \\ V_z \end{bmatrix}",
    )
    para(doc, "Это эталон полёта, а не то, что насчитал БИНС.")

    heading(doc, "2. ДУС и ДЛУ")
    para(
        doc,
        "Внутри модели самолёта датчиков нет. Берём идеал и портим его так, как портит реальный прибор:",
    )
    add_display(
        doc,
        r"\boldsymbol{\omega}_m=(I+K_{\omega})\Phi_{\omega}\boldsymbol{\omega}+\mathbf{b}_{\omega}+\mathbf{n}_{\omega}",
    )
    add_display(
        doc,
        r"\mathbf{a}_m=(I+K_a)\Phi_a\mathbf{a}+\mathbf{b}_a+\mathbf{n}_a",
    )
    bullets(
        doc,
        [
            "$K$ — ошибка масштаба (прибор чуть «длиннее» или «короче» шкалы);",
            "$\\Phi$ — оси датчиков не строго перпендикулярны;",
            "$\\mathbf{b}$ — постоянное смещение нуля, табл. 3: $0.003^{\\circ}/\\mathrm{ч}$ у ДУС, $25\\,\\mu g$ у ДЛУ;",
            "$\\mathbf{n}$ — белый шум.",
        ],
    )
    para(
        doc,
        "В БИНС входят только шесть чисел $\\boldsymbol{\\omega}_m$, $\\mathbf{a}_m$ и память прошлого шага: "
        "$q$, $C_b^n$ и навигационный вектор. Координат самолёта БИНС не знает.",
    )

    heading(doc, "3. ГНСС")
    para(
        doc,
        "Спутники не моделируем. ГНСС — это истина FX1 плюс шум, 10 раз в секунду:",
    )
    add_display(
        doc,
        r"Y_{\mathrm{ГНСС}}=[\varphi,\lambda,h,V_N,V_E,V_h]_{\mathrm{FX1}}+\boldsymbol{\sigma}\circ\mathbf{w},"
        r"\;\mathbf{w}\sim\mathcal{N}(0,I)",
    )
    para(
        doc,
        "СКО из табл. 2: по дальности $6.6$ м (для широты и долготы делим на радиус Земли), "
        "по высоте $6.6$ м, по скорости $0.05$ м/с. Углов и показаний ДУС/ДЛУ у ГНСС нет.",
    )

    heading(doc, "4. БИНС")
    para(
        doc,
        "БИНС сам интегрирует датчики и получает ориентацию, скорость и координаты. "
        "Это обычное счисление в географических осях, не фильтр.",
    )
    heading(doc, "4.1. Ориентация", 2)
    para(
        doc,
        "ДУС измеряет вращение относительно инерциального пространства. Географические оси "
        "(север–восток–вниз) сами поворачиваются: Земля крутится, и самолёт по ней летит. "
        "Чтобы получить вращение самолёта относительно этих осей, из показания ДУС вычитаем "
        "вращение самих осей:",
    )
    add_display(
        doc,
        r"\boldsymbol{\omega}_b^n=\boldsymbol{\omega}_m-\boldsymbol{\omega}_{ni}^b,\ "
        r"\boldsymbol{\omega}_{ni}=\boldsymbol{\omega}_{ie}+\boldsymbol{\omega}_{en}",
    )
    para(
        doc,
        "Здесь $\\boldsymbol{\\omega}_{ie}$ — вращение Земли, $\\boldsymbol{\\omega}_{en}$ — поворот осей "
        "из-за движения над Землёй. Ориентацию ведём кватернионом, без деления на $\\cos\\vartheta$:",
    )
    add_display(
        doc,
        r"\dot{q}=\frac{1}{2}q\otimes\boldsymbol{\omega}_b^n,\ q=[q_0,q_1,q_2,q_3]",
    )
    para(
        doc,
        "Из кватерниона собираем матрицу $C_b^n$. Углы курса, тангажа и крена из неё: "
        "$(\\psi,\\vartheta,\\gamma)=\\mathrm{c\\_ang}(C_b^n)$.",
    )

    heading(doc, "4.2. Матрица ориентации $C_b^n$", 2)
    para(
        doc,
        "Эта матрица переводит вектор из осей самолёта в географические оси. "
        "В MATLAB та же матрица называется $C2$: ось $X$ на север, $Y$ на восток, $Z$ вниз.",
    )
    add_display(
        doc,
        r"\mathbf{v}^n=C_b^n\mathbf{v}^b,\ (C_b^n)^{\top}C_b^n=I,\ \det=+1",
    )
    para(
        doc,
        "Через курс $\\psi$, тангаж $\\vartheta$ и крен $\\gamma$ (как в BINS_OFK_2ch.m):",
    )
    add_display(
        doc,
        r"C_b^n=\begin{bmatrix}"
        r"\cos\vartheta\cos\psi & -\cos\gamma\sin\psi+\sin\gamma\sin\vartheta\cos\psi & \sin\gamma\sin\psi+\cos\gamma\sin\vartheta\cos\psi \\"
        r"\cos\vartheta\sin\psi & \cos\gamma\cos\psi+\sin\gamma\sin\vartheta\sin\psi & -\sin\gamma\cos\psi+\cos\gamma\sin\vartheta\sin\psi \\"
        r"-\sin\vartheta & \sin\gamma\cos\vartheta & \cos\gamma\cos\vartheta"
        r"\end{bmatrix}",
    )
    bullets(
        doc,
        [
            "Строки — север, восток, вниз.",
            "Столбцы — оси самолёта $x$, $y$, $z$.",
            "Элемент $C_{31}=-\\sin\\vartheta$ связан с тангажом.",
        ],
    )
    para(
        doc,
        "Это не датчик, а результат интегрирования ДУС. В ОФК-1 передаём всю матрицу $3\\times 3$, "
        "не три угла по отдельности.",
    )

    heading(doc, "4.3. Скорость и координаты", 2)
    para(
        doc,
        "ДЛУ измеряет удельное ускорение в осях самолёта. Сначала переводим его в географические оси "
        "матрицей $C_b^n$. Потом добавляем тяжесть и поправки, потому что оси не инерциальные: "
        "Земля вращается, и сами оси поворачиваются, пока самолёт летит.",
    )
    add_display(
        doc,
        r"\dot{\mathbf{V}}^n=C_b^n\mathbf{a}_m+\mathbf{g}+\mathbf{a}_{\mathrm{пер}}+\mathbf{a}_{\mathrm{кор}}",
    )
    bullets(
        doc,
        [
            "$C_b^n\\mathbf{a}_m$ — показание ДЛУ, переведённое на север, восток, вниз;",
            "$\\mathbf{g}$ — ускорение свободного падения (в $\\mathbf{a}_m$ его не было);",
            "$\\mathbf{a}_{\\mathrm{кор}}=-2\\boldsymbol{\\omega}_{ie}\\times\\mathbf{V}$ — кориолисово ускорение: "
            "Земля вращается, самолёт относительно неё движется;",
            "$\\mathbf{a}_{\\mathrm{пер}}$ — переносное ускорение. Географические оси поворачиваются, "
            "потому что самолёт летит над Землёй и потому что Земля вращается. В коде это "
            "$-\\boldsymbol{\\omega}_{en}\\times\\mathbf{V}$ и центробежный член "
            "$-\\boldsymbol{\\omega}_{ie}\\times(\\boldsymbol{\\omega}_{ie}\\times\\mathbf{r})$.",
        ],
    )
    para(doc, "Высота, широта и долгота из скорости:")
    add_display(
        doc,
        r"\dot{h}=V_h,\ \dot{\varphi}=\frac{V_N}{R_N},\ "
        r"\dot{\lambda}=\frac{V_E}{R_E\cos\varphi}",
    )

    heading(doc, "4.4. Что выдаёт БИНС", 2)
    bullets(
        doc,
        [
            "навигационный вектор $Y_{\\mathrm{БИНС}}=[V_N,V_h,V_E,h,\\varphi,\\lambda]$ "
            "(такой порядок в программе);",
            "кватернион $q$ и матрица $C_b^n$;",
            "копию $\\mathbf{a}_m$ — её ОФК-1 берёт, когда собирает матрицы фильтра.",
        ],
    )
    para(
        doc,
        "По составу это те же шесть величин, что у ГНСС: широта, долгота, высота и три скорости. "
        "Вертикальную скорость $V_h$ храним вверх. В горизонтальных кусках матрицы $F$ из MATLAB "
        "берут $V_h=-NP(2)$, потому что там ось $Z$ направлена вниз.",
    )

    heading(doc, "5. Невязка $\\varepsilon$")
    para(
        doc,
        "Это единственное измерение ОФК-1: БИНС минус ГНСС. Порядок чисел такой же, как строки матрицы $H$:",
    )
    add_display(
        doc,
        r"\varepsilon=Z=\begin{bmatrix}"
        r"\varphi_{\mathrm{Б}}-\varphi_{\mathrm{Г}} \\"
        r"\lambda_{\mathrm{Б}}-\lambda_{\mathrm{Г}} \\"
        r"V_{N\mathrm{Б}}-V_{N\mathrm{Г}} \\"
        r"V_{E\mathrm{Б}}-V_{E\mathrm{Г}} \\"
        r"V_{h\mathrm{Б}}-V_{h\mathrm{Г}} \\"
        r"h_{\mathrm{Б}}-h_{\mathrm{Г}}"
        r"\end{bmatrix}",
    )
    para(
        doc,
        "Сам вектор $Y_{\\mathrm{БИНС}}$ вторым измерением в фильтр не кладём. "
        "Он нужен, чтобы посчитать эту разность и чтобы собрать матрицу $F$ в текущей точке полёта.",
    )

    heading(doc, "6. ОФК-1")
    para(
        doc,
        "Фильтр Калмана по ошибкам БИНС. Шаг 10 Гц, когда пришёл ГНСС. "
        "Оценки в счисление БИНС не возвращаем: БИНС так и остаётся автономным.",
    )
    heading(doc, "6.1. Что на входе", 2)
    bullets(
        doc,
        [
            "невязка $\\varepsilon$ — измерение;",
            "$\\mathbf{a}_m$ и $C_b^n$ — чтобы собрать матрицы $F$ и $G$;",
            "предыдущие оценка $\\hat{\\mathbf{x}}$ и ковариация $P$;",
            "шаг $dt=0.1$ с и матрица шума измерений $R$ из СКО ГНСС;",
            "текущие скорость, широта и высота БИНС — только чтобы подставить числа в $F$. "
            "Сырой ДУС $\\boldsymbol{\\omega}_m$ в фильтр не подаём: он уже внутри $C_b^n$.",
        ],
    )

    heading(doc, "6.2. Что оцениваем (15 ошибок БИНС, не самолёта)", 2)
    add_display(
        doc,
        r"\mathbf{x}=["
        r"\delta\vartheta_1,\delta\vartheta_2,\delta\vartheta_3,"
        r"\delta V_N,\delta V_E,\delta\varphi,\delta\lambda,"
        r"\Delta a_x,\Delta a_y,\Delta a_z,"
        r"\Delta\omega_x,\Delta\omega_y,\Delta\omega_z,"
        r"\delta V_h,\delta h"
        r"]^{\top}",
    )
    para(
        doc,
        "Первые три — малые углы рассогласования ориентации. Дальше ошибки скорости и координат, "
        "смещения ДЛУ и ДУС, вертикальная скорость и высота. "
        "$\\Delta\\boldsymbol{\\omega}$ — смещение гироскопа: тому, кто дальше берёт угловую скорость, "
        "можно отдать $\\boldsymbol{\\omega}_m-\\Delta\\hat{\\boldsymbol{\\omega}}$. На коротком полёте эта оценка слабая.",
    )

    heading(doc, "6.3. Уравнения фильтра", 2)
    add_display(
        doc,
        r"\dot{\mathbf{x}}=F\mathbf{x}+G\mathbf{w},\ Z=H\mathbf{x}+\mathbf{v}",
    )
    bullets(
        doc,
        [
            "$F$ — матрица $15\\times 15$. Первые 13 строк и столбцов как в MATLAB BINS_OFK_2ch.m: "
            "вращение Земли, кориолис, переносное ускорение, удельное ускорение через $C_b^n$. "
            "Смещения датчиков считаем постоянными (их строки в $F$ нулевые).",
            "Вертикаль дописана из того же счисления: ошибка $\\delta V_h$ зависит от ориентации, "
            "смещения ДЛУ, ошибок горизонтальной скорости, широты и высоты "
            "(член $2|\\mathbf{g}|/R$). Высота: $\\delta\\dot{h}=\\delta V_h$. "
            "Ошибка $\\delta V_h$ чуть входит в горизонтальные ускорения.",
            "$G$ — матрица $15\\times 6$: шум ДУС идёт в ориентацию, шум ДЛУ — в скорость, "
            "оба через $C_b^n$.",
            "$H$ — матрица $6\\times 15$: единицы стоят только на ошибках "
            "$\\delta\\varphi$, $\\delta\\lambda$, $\\delta V_N$, $\\delta V_E$, $\\delta V_h$, $\\delta h$. "
            "Углы и смещения датчиков в измерение не входят.",
            "Шум процесса: $Q=\\mathrm{diag}(W_q^2/dt)$, "
            "$W_q$ из трёх чисел $20\\cdot 10^{-10}$ и трёх чисел $25\\cdot 10^{-7}$ — как в MATLAB.",
            "Шум измерения: $R=\\mathrm{diag}(\\sigma_{\\mathrm{ГНСС}}^2)$ в том же порядке, что $Z$.",
            "Стартовая $P_0$: квадраты $\\sigma$ из табл. 5, для смещений — из табл. 3.",
        ],
    )

    heading(doc, "6.4. Один шаг фильтра", 2)
    para(doc, "Непрерывную модель переводим на шаг ГНСС $\\Delta t=0.1$ с:")
    add_display(
        doc,
        r"\Phi\approx e^{F\Delta t}\approx I+F\Delta t+\frac{(F\Delta t)^2}{2},\ G_d=G\Delta t",
    )
    para(doc, "Сначала прогноз (без ГНСС оценка просто протягивается моделью ошибок):")
    add_display(
        doc,
        r"\hat{\mathbf{x}}^-=\Phi\hat{\mathbf{x}},\ "
        r"P^-=\Phi P\Phi^{\top}+G_d Q G_d^{\top}",
    )
    para(
        doc,
        "Потом поправка по невязке. $P$ считаем по формуле Джозефа, чтобы ковариация оставалась устойчивой:",
    )
    add_display(
        doc,
        r"\nu=Z-H\hat{\mathbf{x}}^-,\ S=HP^-H^{\top}+R,\ "
        r"K=P^-H^{\top}S^{-1}",
    )
    add_display(
        doc,
        r"\hat{\mathbf{x}}=\hat{\mathbf{x}}^-+K\nu,\ "
        r"P=(I-KH)P^-(I-KH)^{\top}+KRK^{\top}",
    )

    heading(doc, "6.5. Что на выходе", 2)
    para(
        doc,
        "Вектор из 15 оценок $\\hat{\\mathbf{x}}$ и матрица $P$ размера $15\\times 15$. "
        "В БИНС их не подставляем: счисление как шло по сырым датчикам, так и идёт.",
    )
    para(
        doc,
        "Тому, кто берёт навигацию дальше (потом ОФК-2), отдаём поправленные скорость и координаты "
        "и угловую скорость без оценённого смещения ДУС:",
    )
    add_display(
        doc,
        r"V_N^c=V_N-\delta\hat{V}_N,\ "
        r"V_E^c=V_E-\delta\hat{V}_E,\ "
        r"V_h^c=V_h-\delta\hat{V}_h",
    )
    add_display(
        doc,
        r"\varphi^c=\varphi-\delta\hat{\varphi},\ "
        r"\lambda^c=\lambda-\delta\hat{\lambda},\ "
        r"h^c=h-\delta\hat{h}",
    )
    add_display(
        doc,
        r"\boldsymbol{\omega}^c=\boldsymbol{\omega}_m-\Delta\hat{\boldsymbol{\omega}}",
    )
    para(
        doc,
        "Матрицу $C_b^n$ не крутим. Смещение ДЛУ $\\Delta\\mathbf{a}$ из показания $\\mathbf{a}_m$ не вычитаем.",
    )

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(16)
    r = note.add_run(
        "ОФК-2 в это письмо не входит. Готов уточнить любой блок по вашей пометке."
    )
    set_run_font(r)
    r.italic = True

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"saved: {path}")
